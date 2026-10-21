"""Card activation: the Card No extractor, delivery/acceptance flow, PDF requests and permissions."""
import shutil
import tempfile

from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase, override_settings
from django.urls import reverse

from io import BytesIO

from openpyxl import load_workbook

from apps.accounts.models import Branch, CsPerm, Role, User
from apps.core.testing import PASSWORD, make_world, xlsx_statement

from . import reports
from .models import ActivationRequest, Card, CardBlock, CardStatus, FollowUp, RejectReason, RequestStatus
from .services import (CardError, accept_cards, block_cards, dashboard, deliver_cards, import_card_file, read_card_numbers, resubmit_request, submit_request,
                       suggest_cards, update_request_status)

BIN = "461699"
PDF = b"%PDF-1.4\n1 0 obj<<>>endobj\ntrailer<<>>\n%%EOF"
CARD_ROWS = [["Card No", "Date Received", "Date Disbursed"], [4000000000000003, "x", ""], ["4000000000000001", "", ""], [4000000000000002.0, "", ""],
             ["4000000000000001", "", ""], [None, "", ""]]


def pdf(name="ALI BIN ABU.pdf", content=PDF):
    return SimpleUploadedFile(name, content, content_type="application/pdf")


class ExtractorTests(TestCase):
    def test_only_card_no_in_file_order_without_sorting(self):
        numbers = read_card_numbers(xlsx_statement(CARD_ROWS), "cards.xlsx")
        self.assertEqual(numbers, ["4000000000000003", "4000000000000001", "4000000000000002", "4000000000000001"])

    def test_csv_and_header_with_other_spelling(self):
        self.assertEqual(read_card_numbers(b"Other,CARD NO.\nx,111\ny,222\n", "c.csv"), ["111", "222"])

    def test_missing_column_and_bad_extension(self):
        with self.assertRaises(CardError):
            read_card_numbers(b"A,B\n1,2\n", "c.csv")
        with self.assertRaises(CardError):
            read_card_numbers(b"x", "c.txt")


class CardFlowTests(TestCase):
    def setUp(self):
        self.w = make_world()
        self.cs = User.objects.create_user("cs@t.test", PASSWORD, name="CS", role=Role.CS)

    def test_import_keeps_order_skips_known_and_repeated_numbers(self):
        result = import_card_file(self.cs, xlsx_statement(CARD_ROWS), "cards.xlsx")
        self.assertEqual((result["added"], result["repeated"]), (3, 1))
        self.assertEqual(list(Card.objects.values_list("card_no", flat=True)), ["4000000000000003", "4000000000000001", "4000000000000002"])
        again = import_card_file(self.cs, xlsx_statement(CARD_ROWS), "cards2.xlsx")
        self.assertEqual(again["added"], 0)
        self.assertEqual(again["already_stored"], 4)

    def test_deliver_then_accept_by_the_destination_branch_only(self):
        import_card_file(self.cs, xlsx_statement(CARD_ROWS), "cards.xlsx")
        ids = list(Card.objects.values_list("pk", flat=True))[:2]
        self.assertEqual(deliver_cards(self.cs, ids, self.w["kj"]), 2)
        card = Card.objects.get(pk=ids[0])
        self.assertEqual((card.status, card.branch), (CardStatus.DELIVERED, self.w["kj"]))
        self.assertTrue(self.w["teller"].notifications.exists())
        with self.assertRaises(CardError):
            accept_cards(self.w["teller2"], ids)              # another branch
        self.assertEqual(accept_cards(self.w["teller"], ids), 2)
        card.refresh_from_db()
        self.assertEqual(card.status, CardStatus.ACCEPTED)
        self.assertIsNotNone(card.accepted_at)
        with self.assertRaises(CardError):
            deliver_cards(self.cs, ids, self.w["kj"])         # no longer Not Sent


def batch_numbers(count=6, start=1):
    """A card file with `count` card numbers that start with the BIN (4616990000000001 ...)."""
    return xlsx_statement([["Card No"]] + [[f"{BIN}{n:010d}"] for n in range(start, start + count)])


def make_cs(email, rights):
    return User.objects.create_user(email, PASSWORD, name=email, role=Role.CS, cs_permissions=rights)


def customer(**over):
    data = {"customer_name": "Siti Aminah", "account_no": "1150011419969", "id_no": "3521166602990001", "phone": "0187855589", "reason": 1}
    data.update(over)
    return data


class DashboardTests(TestCase):
    """Dashboard numbers follow every action: new batch, delivery, card request, block."""

    def setUp(self):
        self.w = make_world()
        self.cs = make_cs("cs@t.test", CsPerm.BATCH | CsPerm.ACTIVATION | CsPerm.BLOCK)

    def figures(self, branch=None):
        d = dashboard(self.w["teller"])
        return d["total"] if branch is None else next(b for b in d["branches"] if b["id"] == branch.pk)

    def test_batch_delivery_request_and_block_move_the_numbers(self):
        import_card_file(self.cs, batch_numbers(6), "b.xlsx")
        self.assertEqual(self.figures()["undisbursed"], 6)
        ids = list(Card.objects.values_list("pk", flat=True))
        deliver_cards(self.cs, ids[:4], self.w["kj"])        # head office -> Kajang
        total, kj = self.figures(), self.figures(self.w["kj"])
        self.assertEqual((total["undisbursed"], total["disbursed"], total["remaining"], total["collected"]), (2, 4, 4, 0))
        self.assertEqual((kj["disbursed"], kj["remaining"], kj["collected"], kj["blocked"]), (4, 4, 0, 0))
        submit_request(self.w["teller"], customer(), pdf(), f"{BIN}0000000001")             # a customer makes a card
        kj = self.figures(self.w["kj"])
        self.assertEqual((kj["disbursed"], kj["remaining"], kj["collected"]), (4, 3, 1))
        block_cards(self.cs, [ids[1]], "Lost in transit")                                   # CS blocks one
        kj = self.figures(self.w["kj"])
        self.assertEqual((kj["remaining"], kj["blocked"], kj["collected"], kj["disbursed"]), (2, 1, 1, 4))
        self.assertEqual(self.figures()["blocked"], 1)

    def test_branch_without_cards_shows_a_dash_and_a_new_branch_is_added_automatically(self):
        self.assertTrue(self.figures(self.w["pj"])["empty"])
        new = Branch.objects.create(code="SA", name="Shah Alam")
        data = dashboard(self.w["teller"])
        self.assertIn("Shah Alam", [b["label"] for b in data["branches"]])
        self.assertTrue(next(b for b in data["branches"] if b["id"] == new.pk)["empty"])
        self.client.login(username="teller@t.test", password=PASSWORD)
        html = self.client.get(reverse("cards:dashboard")).content.decode()
        self.assertIn("Shah Alam", html)
        self.assertRegex(html, r'data-m="b%d\.remaining"[^>]*>-<' % new.pk)          # "Remaining Card: -"

    def test_own_branch_comes_first_and_everybody_sees_all_branches(self):
        d = dashboard(self.w["teller2"])
        self.assertEqual([(b["code"], b["mine"]) for b in d["branches"]], [("PJ", True), ("KJ", False)])    # own branch first
        self.assertEqual([b["code"] for b in dashboard(self.w["teller"])["branches"]], ["KJ", "PJ"])
        self.assertFalse(any(b["mine"] for b in dashboard(self.cs)["branches"]))
        self.assertEqual(len(dashboard(self.cs)["branches"]), 2)

    def test_dashboard_is_its_own_page_and_json_matches(self):
        import_card_file(self.cs, batch_numbers(3), "b.xlsx")
        self.client.login(username="teller@t.test", password=PASSWORD)
        html = self.client.get(reverse("cards:dashboard")).content.decode()
        for label in ("Total Undisbursed Card", "Total Disbursed", "Collected by Customer", "Remaining Card", "Block Card"):
            self.assertIn(label, html)
        self.assertNotIn("Card dashboard", self.client.get(reverse("cards:teller_requests")).content.decode())   # its own page now
        data = self.client.get(reverse("cards:dashboard_data")).json()
        self.assertEqual(data["total"]["undisbursed"], 3)
        self.client.logout()
        self.client.login(username="finance@t.test", password=PASSWORD)
        self.assertEqual(self.client.get(reverse("cards:dashboard_data")).status_code, 403)


class CardRequestTests(TestCase):
    def setUp(self):
        self.w = make_world()
        self.cs = make_cs("cs@t.test", CsPerm.DEFAULT)
        import_card_file(self.cs, batch_numbers(6), "b.xlsx")
        self.ids = list(Card.objects.values_list("pk", flat=True))
        deliver_cards(self.cs, self.ids[:3], self.w["kj"])
        deliver_cards(self.cs, self.ids[3:4], self.w["pj"])

    def test_suggestions_are_the_own_branch_cards_most_similar_first(self):
        self.assertEqual(suggest_cards(self.w["teller"], ""), [f"{BIN}000000000{n}" for n in (1, 2, 3)])
        self.assertEqual(suggest_cards(self.w["teller"], "0000000002"), [f"{BIN}0000000002"])
        self.assertEqual(suggest_cards(self.w["teller"], "2"), [f"{BIN}0000000002"])     # also found by a part inside the number
        self.assertEqual(suggest_cards(self.w["teller"], BIN + "000000000"), [f"{BIN}000000000{n}" for n in (1, 2, 3)])   # typed with the BIN: same
        self.assertEqual(suggest_cards(self.w["teller"], "4"), [])                      # card 4 belongs to PJ
        self.assertEqual(suggest_cards(self.cs, ""), [])                                # only Tellers have a branch

    def test_request_takes_the_card_and_keeps_the_customer_data(self):
        req = submit_request(self.w["teller"], customer(), pdf(), f"{BIN}0000000002")
        card = Card.objects.get(card_no=f"{BIN}0000000002")
        self.assertEqual((card.status, req.card_id, req.status, req.customer_name, req.reason), (CardStatus.COLLECTED, card.pk, RequestStatus.SENT, "SITI AMINAH", 1))
        self.assertIsNotNone(card.collected_at)
        self.assertIsNotNone(card.accepted_at)                    # handed over from Delivered: receipt confirmed at the same moment
        self.assertEqual(suggest_cards(self.w["teller"], "2"), [])  # no longer offered

    def test_card_must_be_free_and_of_the_own_branch(self):
        for number, why in ((f"{BIN}0000000004", "other branch"), (f"{BIN}0000000005", "not delivered"), (f"{BIN}9999999999", "unknown"),
                            ("400000000000001", "wrong BIN")):
            with self.assertRaises(CardError, msg=why):
                submit_request(self.w["teller"], customer(), pdf(), number)
        submit_request(self.w["teller"], customer(), pdf(), f"{BIN}0000000001")
        with self.assertRaises(CardError):
            submit_request(self.w["teller"], customer(), pdf(), f"{BIN}0000000001")           # used twice
        self.assertEqual(ActivationRequest.objects.count(), 1)

    def test_customer_fields_are_checked_and_nothing_is_saved_on_error(self):
        for bad in ({"customer_name": ""}, {"customer_name": "N2iqa"}, {"account_no": "12AB"}, {"id_no": "12"}, {"phone": "abc"}, {"reason": 9}, {"reason": ""}):
            with self.assertRaises(CardError, msg=bad):
                submit_request(self.w["teller"], customer(**bad), pdf(), f"{BIN}0000000001")
        for bad_pdf in (pdf("x.pdf", b"not a pdf"), pdf("x.txt")):
            with self.assertRaises(CardError):
                submit_request(self.w["teller"], customer(), bad_pdf, f"{BIN}0000000001")
        with self.assertRaises(CardError):
            submit_request(self.w["teller"], customer(), None, f"{BIN}0000000001")
        self.assertEqual(ActivationRequest.objects.count(), 0)
        self.assertEqual(Card.objects.get(card_no=f"{BIN}0000000001").status, CardStatus.DELIVERED)

    def test_status_flow_and_dates(self):
        req = submit_request(self.w["teller"], customer(), pdf(), f"{BIN}0000000001")
        with self.assertRaises(CardError):
            update_request_status(self.cs, req.pk, RequestStatus.ACTIVATED)     # must be requested first
        update_request_status(self.cs, req.pk, RequestStatus.ACTIVATION_REQUESTED)
        update_request_status(self.cs, req.pk, RequestStatus.ACTIVATED)
        req.refresh_from_db()
        self.assertEqual(req.status, RequestStatus.ACTIVATED)
        self.assertIsNotNone(req.mabes_at)
        self.assertIsNotNone(req.activated_at)

    def test_reject_needs_a_type_and_a_reason_from_the_lists(self):
        req = submit_request(self.w["teller"], customer(), pdf(), f"{BIN}0000000001")
        for follow_up, reason in ((None, None), (FollowUp.RETURN, None), (None, RejectReason.KTP), (9, 1), (1, 9)):
            with self.assertRaises(CardError, msg=(follow_up, reason)):
                update_request_status(self.cs, req.pk, RequestStatus.REJECTED, follow_up, reason)
        update_request_status(self.cs, req.pk, RequestStatus.REJECTED, FollowUp.INCOMPLETE, RejectReason.MOTHER_NAME)
        req.refresh_from_db()
        self.assertEqual((req.get_follow_up_display(), req.note_text), ("Incomplete Data", "Nama Ibu Kandung"))

    def test_incomplete_and_return_can_be_sent_again_close_cannot(self):
        for number, follow_up, again in ((1, FollowUp.INCOMPLETE, True), (2, FollowUp.RETURN, True), (3, FollowUp.CLOSE, False)):
            req = submit_request(self.w["teller"], customer(customer_name=f"Name {number}".replace("1", "A").replace("2", "B").replace("3", "C")), pdf(), f"{BIN}000000000{number}")
            update_request_status(self.cs, req.pk, RequestStatus.REJECTED, follow_up, RejectReason.SIGNATURE)
            if again:
                resubmit_request(self.w["teller"], req.pk, pdf("v2.pdf"))
                req.refresh_from_db()
                self.assertEqual((req.status, req.follow_up, req.reject_reason, req.resubmit_count), (RequestStatus.SENT, None, None, 1))
            else:
                with self.assertRaises(CardError):
                    resubmit_request(self.w["teller"], req.pk, pdf("v2.pdf"))

    def test_pdf_is_served_inline_to_cs_and_own_branch_only(self):
        req = submit_request(self.w["teller"], customer(), pdf(), f"{BIN}0000000001")
        url = reverse("cards:request_file", args=[req.pk])
        self.client.login(username="cs@t.test", password=PASSWORD)
        response = self.client.get(url)
        self.assertEqual((response.status_code, response["Content-Type"]), (200, "application/pdf"))
        self.assertIn("inline", response["Content-Disposition"])
        self.client.logout()
        self.client.login(username="teller2@t.test", password=PASSWORD)
        self.assertEqual(self.client.get(url).status_code, 403)

    def test_form_page_posts_and_suggest_endpoint(self):
        self.client.login(username="teller@t.test", password=PASSWORD)
        body = self.client.get(reverse("cards:card_suggest"), {"q": "3"}).json()
        self.assertEqual(body, {"bin": BIN, "suggestions": ["0000000003"]})
        response = self.client.post(reverse("cards:teller_submit"), {**customer(), "card_suffix": "0000000003", "file": pdf()})
        self.assertRedirects(response, reverse("cards:teller_requests"))
        self.assertEqual(Card.objects.get(card_no=f"{BIN}0000000003").status, CardStatus.COLLECTED)
        wrong = self.client.post(reverse("cards:teller_submit"), {**customer(), "card_suffix": "0000000003", "file": pdf()})
        self.assertEqual(wrong.status_code, 422)                                         # used already: the form is shown again with the message
        self.assertContains(wrong, "not available at your branch", status_code=422)


class PermissionTests(TestCase):
    """CS A: activation only. CS B: activation + block. CS C: batch only."""

    def setUp(self):
        self.w = make_world()
        self.a = make_cs("a@t.test", CsPerm.ACTIVATION)
        self.b = make_cs("b@t.test", CsPerm.ACTIVATION | CsPerm.BLOCK)
        self.c = make_cs("c@t.test", CsPerm.BATCH)
        import_card_file(self.c, batch_numbers(4), "b.xlsx")
        self.ids = list(Card.objects.values_list("pk", flat=True))
        deliver_cards(self.c, self.ids, self.w["kj"])
        self.req = submit_request(self.w["teller"], customer(), pdf(), f"{BIN}0000000001")

    def post(self, email, name, data=None, args=()):
        self.client.logout()
        self.client.login(username=email, password=PASSWORD)
        return self.client.post(reverse(name, args=args), data or {})

    def test_block_only_for_cs_with_the_right(self):
        self.assertEqual(self.post("a@t.test", "cards:cs_block", {"card": self.ids[1], "note": "lost"}).status_code, 403)
        self.assertEqual(Card.objects.get(pk=self.ids[1]).status, CardStatus.DELIVERED)
        self.assertEqual(self.post("b@t.test", "cards:cs_block", {"card": self.ids[1], "note": "lost"}).status_code, 302)
        self.assertEqual(Card.objects.get(pk=self.ids[1]).status, CardStatus.BLOCKED)
        self.assertEqual(CardBlock.objects.get(card_id=self.ids[1]).note, "lost")
        with self.assertRaises(CardError):
            block_cards(self.a, [self.ids[2]], "x")                 # the service refuses too
        with self.assertRaises(CardError):
            block_cards(self.b, [self.ids[2]], "  ")                # a reason is required
        with self.assertRaises(CardError):
            block_cards(self.b, [self.ids[0]], "customer card")     # already collected by a customer: cannot be blocked

    def test_activation_only_for_cs_with_the_right(self):
        url = ("cards:cs_request_status", (self.req.pk,))
        self.assertEqual(self.post("c@t.test", url[0], {"status": 1}, url[1]).status_code, 403)
        self.assertEqual(self.post("a@t.test", url[0], {"status": 1}, url[1]).status_code, 302)
        self.req.refresh_from_db()
        self.assertEqual(self.req.status, RequestStatus.ACTIVATION_REQUESTED)

    def test_batch_and_deliver_only_for_cs_with_the_right(self):
        upload = SimpleUploadedFile("c.xlsx", batch_numbers(2, start=50))
        self.client.login(username="a@t.test", password=PASSWORD)
        self.assertEqual(self.client.post(reverse("cards:cs_upload"), {"file": upload}).status_code, 403)
        self.assertEqual(self.post("b@t.test", "cards:cs_deliver", {"branch": self.w["kj"].pk}).status_code, 403)
        with self.assertRaises(CardError):
            import_card_file(self.b, batch_numbers(2, start=50), "c.xlsx")

    def test_screens_show_only_what_the_cs_may_do(self):
        self.client.login(username="a@t.test", password=PASSWORD)
        html = self.client.get(reverse("cards:cs_cards") + "?tab=delivered").content.decode()
        self.assertNotIn("Block selected", html)
        self.assertNotIn("Add card batch", html)
        self.client.logout()
        self.client.login(username="b@t.test", password=PASSWORD)
        self.assertIn("Block selected", self.client.get(reverse("cards:cs_cards") + "?tab=delivered").content.decode())
        self.assertIn("Update status", self.client.get(reverse("cards:cs_requests")).content.decode())

    def test_admin_sets_the_rights_in_the_user_form(self):
        admin_client = self.client
        admin_client.login(username="admin@t.test", password=PASSWORD)
        response = admin_client.post(reverse("accounts:user_edit", args=[self.a.pk]), {
            "name": "CS A", "email": "a@t.test", "role": Role.CS, "branch": "", "cs_rights": [str(CsPerm.ACTIVATION), str(CsPerm.BLOCK)]})
        self.assertEqual(response.status_code, 302)
        self.a.refresh_from_db()
        self.assertEqual(self.a.cs_permissions, CsPerm.ACTIVATION | CsPerm.BLOCK)
        admin_client.post(reverse("accounts:user_edit", args=[self.a.pk]), {"name": "CS A", "email": "a@t.test", "role": Role.FINANCE, "branch": ""})
        self.a.refresh_from_db()
        self.assertEqual(self.a.cs_permissions, 0)                 # not a CS any more: no rights kept


class ReportTests(TestCase):
    def setUp(self):
        self.w = make_world()
        self.cs = make_cs("cs@t.test", CsPerm.BATCH | CsPerm.ACTIVATION | CsPerm.BLOCK)
        import_card_file(self.cs, batch_numbers(6), "b.xlsx")
        ids = list(Card.objects.values_list("pk", flat=True))
        deliver_cards(self.cs, ids[:5], self.w["kj"])
        self.done = submit_request(self.w["teller"], customer(), pdf(), f"{BIN}0000000001")
        update_request_status(self.cs, self.done.pk, RequestStatus.ACTIVATION_REQUESTED)
        update_request_status(self.cs, self.done.pk, RequestStatus.ACTIVATED)
        self.bad = submit_request(self.w["teller"], customer(customer_name="Budi", reason=2), pdf(), f"{BIN}0000000002")
        update_request_status(self.cs, self.bad.pk, RequestStatus.REJECTED, FollowUp.INCOMPLETE, RejectReason.BIRTH)
        block_cards(self.cs, [ids[2]], "Damaged")

    def sheet(self, built, name=None):
        data, filename = built
        wb = load_workbook(BytesIO(data))
        return wb, (wb[name] if name else wb.worksheets[0]), filename

    def test_teller_report_matches_the_teller_file(self):
        wb, ws, name = self.sheet(reports.teller_report(self.w["teller"]))
        self.assertRegex(name, r"^Card - Kajang \d+_\d+_\d{4}\.xlsx$")
        self.assertEqual(wb.sheetnames, ["Card Data - Kajang"])
        self.assertEqual([c.value for c in ws[1]], reports.HEADERS)
        head = ws["A1"]
        self.assertEqual((head.fill.fgColor.rgb, head.font.name, head.font.b, head.alignment.horizontal, head.border.left.style), ("FFA4C2F4", "Arial", True, "center", "thin"))
        self.assertEqual((ws.freeze_panes, ws["B2"].number_format, ws["A2"].number_format, ws["M2"].fill.fgColor.rgb), ("A2", "dd/MM/yyyy", "@", "FFFFFFFF"))
        self.assertEqual((ws.column_dimensions["A"].width, ws.column_dimensions["J"].width, ws.column_dimensions["M"].width), (18.9, 21.4, 18.9))
        rows = {r[0].value: [c.value for c in r] for r in ws.iter_rows(min_row=2) if r[0].value}
        self.assertEqual(len(rows), 5)                         # only the 5 cards of Kajang, none of head office
        done = rows[f"{BIN}0000000001"]
        self.assertEqual(done[3], "Kajang")
        self.assertEqual((done[9], done[10], done[12], done[13], done[14], done[15], done[16], done[17]),
                         ("Activation Complete", None, "SITI AMINAH", "3521166602990001", "1150011419969", "0187855589", "Terbit Kartu Baru", "Done"))
        bad = rows[f"{BIN}0000000002"]
        self.assertEqual((bad[9], bad[10], bad[11], bad[16]), ("Pending Activation", "Incomplete Data", "TTL Berbeda", "Ganti Kartu"))
        self.assertEqual(rows[f"{BIN}0000000003"][9:12], ["Blocked", None, "Damaged"])
        self.assertEqual(rows[f"{BIN}0000000004"][9:], [None] * 9)            # nothing happened to this card yet

    def test_red_rows_and_dropdowns_are_in_the_file(self):
        _wb, ws, _ = self.sheet(reports.teller_report(self.w["teller"]))
        rules = [(str(rng.sqref), r) for rng, rs in ws.conditional_formatting._cf_rules.items() for r in rs]
        self.assertEqual([(s, r.formula, r.dxf.fill.bgColor.rgb) for s, r in rules], [("A2:R500", ['$K2<>""'], "FFF4CCCC")])
        lists = {str(dv.sqref): dv.formula1 for dv in ws.data_validations.dataValidation if dv.type == "list"}
        self.assertEqual(lists["K2:K500"], '"Incomplete Data,Return,Close"')
        self.assertEqual(lists["L2:L500"], '"Nama Ibu Kandung,Tanda Tangan Berbeda,TTL Berbeda,KTP Tidak Sesuai DUKCAPIL"')
        self.assertEqual(lists["Q2:Q500"], '"Terbit Kartu Baru,Ganti Kartu,Kad Expired"')
        self.assertEqual(lists["J2:J500"], '"Pending Activation,Activation Complete"')

    def test_cs_report_has_master_data_and_card_stock(self):
        wb, ws, name = self.sheet(reports.cs_report(self.cs), "Master Data")
        self.assertRegex(name, r"^Card Monitoring - HQ \d+_\d+_\d{4}\.xlsx$")
        self.assertEqual(wb.sheetnames, ["Master Data", "Card Stock"])
        self.assertEqual(sum(1 for r in ws.iter_rows(min_row=2) if r[0].value), 6)       # every card, also the one at head office
        self.assertEqual((ws["A2"].font.name, ws["B2"].number_format, ws.column_dimensions["M"].width, ws.column_dimensions["N"].width), ("Calibri", "dd/mm/yyyy", 31.4, 22.6))
        stock = wb["Card Stock"]
        self.assertEqual([c.value for c in stock[1]], ["Branch", "Total Disbursed", "Collected by Customer", "Remaining", "Others Income", "Others Income (Money Received)"])
        by_branch = {r[0].value: [c.value for c in r][1:] for r in stock.iter_rows(min_row=2, max_row=3)}
        self.assertEqual(by_branch["Kajang"], [5, 2, 2, 60, 60])                           # 5 disbursed, 2 collected, 2 left (1 blocked), RM30 each
        self.assertEqual(by_branch["Petaling Jaya"], [0, 0, 0, 0, 0])
        self.assertEqual([c.value for c in stock[4]], [None] * 6)
        self.assertEqual([c.value for c in stock[5]][:6], ["Total Undisbursed Card", "Total All Disbursed Card", "Total Collected Card by Customer", "Total Remaining Card",
                                                           "Total Others Income", "Total Others Income (Money Received)"])
        self.assertEqual([c.value for c in stock[6]], [1, 5, 2, 2, 60, 60])
        self.assertEqual((stock["A1"].fill.fgColor.rgb, stock["E2"].number_format), ("FFD2E3FC", '"RM "#,##0.00'))
        self.assertEqual(sum(1 for rs in stock.conditional_formatting._cf_rules.values() for _ in rs), 1)      # Remaining < 10 turns red

    def test_report_download_for_each_role(self):
        for email, prefix in (("teller@t.test", "Card - Kajang"), ("cs@t.test", "Card Monitoring - HQ"), ("admin@t.test", "Card Monitoring - HQ")):
            self.client.logout()
            self.client.login(username=email, password=PASSWORD)
            response = self.client.get(reverse("cards:report"))
            self.assertEqual(response.status_code, 200, email)
            self.assertIn(prefix, response["Content-Disposition"])
            self.assertEqual(load_workbook(BytesIO(response.content)).sheetnames[0], "Card Data - Kajang" if prefix == "Card - Kajang" else "Master Data")
        self.client.logout()
        self.client.login(username="finance@t.test", password=PASSWORD)
        self.assertEqual(self.client.get(reverse("cards:report")).status_code, 403)


