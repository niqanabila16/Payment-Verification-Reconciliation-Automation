"""
Excel reports of the cards, laid out like the files of the business (colours, fonts, borders, column widths, drop-downs, red rows):

  Teller  "Card - <Branch> d_m_yyyy.xlsx"           1 sheet  "Card Data - <Branch>"   the cards of the Teller's branch
  CS      "Card Monitoring - HQ d_m_yyyy.xlsx"      2 sheets "Master Data" (every card) + "Card Stock" (per branch + totals)

Both Master / Data sheets have the same 18 columns. A row turns red (Follow Up filled in) through a conditional format, so it also
turns red when someone types a Follow Up in Excel later. Values are written as values (no formulas), so previewers show them too.
"""
import io
from datetime import date

from django.utils import timezone
from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

from .models import Card, CardReason, CardStatus, FollowUp, RejectReason, RequestStatus
from .services import dashboard

HEADERS = ["Card No", "Date Received", "Date Disbursed", "Destination Branch", "Date Received at Destination", "Customer Collection Date",
           "Date Submit to HQ", "Date Submit to Mabes", "Activation Date", "Current Status", "Follow Up", "Note", "Nama Pemohon", "No ID",
           "No Rekening", "No Phone Nasabah", "Alasan Ganti Kad", "Collection Status"]
STATUS_PENDING, STATUS_COMPLETE, STATUS_BLOCKED = "Pending Activation", "Activation Complete", "Blocked"
DATE_COLUMNS = (2, 3, 5, 6, 7, 8, 9)               # B C E F G H I

BLACK = Side(style="thin", color="FF000000")
BOX = Border(left=BLACK, right=BLACK, top=BLACK, bottom=BLACK)
HEAD_FILL = PatternFill("solid", fgColor="FFA4C2F4")
STOCK_FILL = PatternFill("solid", fgColor="FFD2E3FC")
WHITE_FILL = PatternFill("solid", fgColor="FFFFFFFF")
RED_ROW = PatternFill(start_color="FFF4CCCC", end_color="FFF4CCCC", fill_type="solid")
LOW_STOCK = PatternFill(start_color="FFF28B82", end_color="FFF28B82", fill_type="solid")
ARIAL10 = Font(name="Arial", size=10)
ARIAL10_BOLD = Font(name="Arial", size=10, bold=True)
CALIBRI11 = Font(name="Calibri", size=11, color="FF000000")


def local_date(value) -> date | None:
    return timezone.localtime(value).date() if value else None


def card_rows(branch=None):
    """The data of the 18 columns for every card (a branch, or all cards), in the file order of the uploads."""
    qs = Card.objects.select_related("branch", "block", "request").order_by("upload_id", "position", "id")
    qs = qs.filter(branch=branch).exclude(status=CardStatus.NOT_SENT) if branch is not None else qs
    for card in qs:
        req = getattr(card, "request", None)
        block = getattr(card, "block", None)
        status = None
        if req is not None:
            status = STATUS_COMPLETE if req.status == RequestStatus.ACTIVATED else STATUS_PENDING
        elif card.status == CardStatus.BLOCKED:
            status = STATUS_BLOCKED
        rejected = req is not None and req.status == RequestStatus.REJECTED
        note = (req.get_reject_reason_display() if rejected and req.reject_reason else None) or (block.note if block else None)
        yield [
            card.card_no, local_date(card.created_at), local_date(card.delivered_at), (card.branch.name or card.branch.code) if card.branch else None,
            local_date(card.accepted_at), local_date(card.collected_at),
            local_date(req.submitted_at) if req else None, local_date(req.mabes_at) if req else None, local_date(req.activated_at) if req else None,
            status, req.get_follow_up_display() if rejected and req.follow_up else None, note,
            (req.customer_name or None) if req else None, (req.id_no or None) if req else None, (req.account_no or None) if req else None,
            (req.phone or None) if req else None, req.get_reason_display() if req and req.reason else None,
            "Done" if card.status == CardStatus.COLLECTED else None,
        ]


def _style_data_sheet(ws, rows: list, capacity: int, *, hq: bool) -> None:
    """Header + data rows + empty formatted rows up to `capacity`, widths, freeze, filter, drop-downs and the red-row rule."""
    date_fmt = "dd/mm/yyyy" if hq else "dd/MM/yyyy"
    ws.sheet_format.defaultRowHeight, ws.sheet_format.defaultColWidth = 15.75, 12.63
    for col, head in enumerate(HEADERS, start=1):
        cell = ws.cell(row=1, column=col, value=head)
        cell.font, cell.fill, cell.border = ARIAL10_BOLD, HEAD_FILL, BOX
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    left_wrap = Alignment(horizontal="left", vertical="center", wrap_text=True)
    for index in range(2, capacity + 2):
        values = rows[index - 2] if index - 2 < len(rows) else [None] * len(HEADERS)
        for col, value in enumerate(values, start=1):
            cell = ws.cell(row=index, column=col, value=value)
            cell.border = BOX
            if col == 1:
                cell.number_format = "@"
                cell.font = CALIBRI11 if hq else ARIAL10
                cell.alignment = Alignment(horizontal=None) if hq else left_wrap
            elif col in DATE_COLUMNS:
                cell.number_format, cell.font, cell.alignment = date_fmt, ARIAL10, left_wrap
            elif col in (10, 11, 12):
                cell.number_format, cell.font, cell.alignment = "@", ARIAL10, left_wrap
            elif 13 <= col <= 16 and not hq:        # the Teller file keeps the white, centred Calibri block of the customer columns
                cell.font, cell.fill, cell.alignment = Font(name="Calibri", size=11, color="FF000000"), WHITE_FILL, Alignment(horizontal="center")
            elif 13 <= col <= 16:
                cell.font, cell.alignment = ARIAL10, Alignment(horizontal="left")
                cell.number_format = "@" if col in (14, 15, 16) else "General"      # ID, account and phone are text: keep leading zeros
            else:
                cell.font, cell.alignment = ARIAL10, left_wrap if not hq or col == 4 else Alignment(horizontal="left")
            if col in (14, 15, 16) and not hq:
                cell.number_format = "@"
    for letter, width in (("A", 18.9), ("J", 21.4), ("M", 31.4 if hq else 18.9)) + ((("N", 22.6), ("O", 18.9)) if hq else ()):
        ws.column_dimensions[letter].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:R{max(len(rows) + 1, 2)}"
    last = capacity + 1
    for column, items in (("L", RejectReason.labels), ("R", ["Done"]), ("K", FollowUp.labels), ("Q", CardReason.labels),
                          ("J", [STATUS_PENDING, STATUS_COMPLETE])):
        dv = DataValidation(type="list", formula1='"' + ",".join(items) + '"', allow_blank=True)
        dv.add(f"{column}2:{column}{last}")
        ws.add_data_validation(dv)
    date_dv = DataValidation(type="custom", allow_blank=True, showErrorMessage=True,
                             formula1='OR(NOT(ISERROR(DATEVALUE(B2))), AND(ISNUMBER(B2), LEFT(CELL("format", B2))="D"))')
    date_dv.add(f"B2:C{last}")
    date_dv.add(f"E2:I{last}")
    ws.add_data_validation(date_dv)
    ws.conditional_formatting.add(f"A2:R{last}", FormulaRule(formula=['$K2<>""'], fill=RED_ROW))


def _stock_sheet(ws, user) -> None:
    """Sheet "Card Stock": the numbers of the dashboard per branch, then the totals of the whole company."""
    data = dashboard(user)
    ws.sheet_format.defaultRowHeight, ws.sheet_format.defaultColWidth = 15.75, 12.63
    ws.column_dimensions["A"].width = 21.4

    def put(row, values, *, bold=False, fill=None):
        for col, value in enumerate(values, start=1):
            cell = ws.cell(row=row, column=col, value=value)
            cell.font = ARIAL10_BOLD if bold else ARIAL10
            cell.alignment, cell.border = Alignment(horizontal="center", vertical="center", wrap_text=True), BOX
            if fill:
                cell.fill = fill

    put(1, ["Branch", "Total Disbursed", "Collected by Customer", "Remaining"], bold=True, fill=STOCK_FILL)
    row = 1
    for branch in sorted(data["branches"], key=lambda b: b["label"]):
        row += 1
        put(row, [branch["label"], branch["disbursed"], branch["collected"], branch["remaining"]])
    first, last_branch = 2, row
    put(row + 1, [None] * 4, bold=True, fill=STOCK_FILL)
    put(row + 2, ["Total Undisbursed Card", "Total All Disbursed Card", "Total Collected Card by Customer", "Total Remaining Card"],
        bold=True, fill=STOCK_FILL)
    total = data["total"]
    put(row + 3, [total["undisbursed"], total["disbursed"], total["collected"], total["remaining"]], bold=True, fill=STOCK_FILL)
    if last_branch >= first:
        ws.conditional_formatting.add(f"D{first}:D{last_branch}", CellIsRule(operator="lessThan", formula=["10"], fill=LOW_STOCK))


def stamp(today: date | None = None) -> str:
    today = today or timezone.localdate()
    return f"{today.day}_{today.month}_{today.year}"          # 31_3_2026, like the business files


def teller_report(user) -> tuple[bytes, str]:
    """The Teller's file: only the cards of the Teller's own branch."""
    branch = user.branch
    label = branch.name or branch.code
    rows = list(card_rows(branch))
    wb = Workbook()
    ws = wb.active
    ws.title = f"Card Data - {label}"[:31]
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    _style_data_sheet(ws, rows, max(500, len(rows) + 100) - 1, hq=False)
    return _save(wb), f"Card - {label} {stamp()}.xlsx"


def cs_report(user) -> tuple[bytes, str]:
    """The CS file: every card of every branch (and the cards still at head office) + the stock per branch."""
    rows = list(card_rows())
    wb = Workbook()
    ws = wb.active
    ws.title = "Master Data"
    _style_data_sheet(ws, rows, max(1500, len(rows) + 100), hq=True)
    _stock_sheet(wb.create_sheet("Card Stock"), user)
    return _save(wb), f"Card Monitoring - HQ {stamp()}.xlsx"


def _save(wb) -> bytes:
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
