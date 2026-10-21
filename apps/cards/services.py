"""Card activation rules: reading the card file, delivering / accepting cards, activation requests."""
import csv
import hashlib
import io
import re
from collections import defaultdict
from decimal import Decimal

from django.conf import settings
from django.db import transaction
from django.db.models import Count
from django.db.models.functions import Substr
from django.utils import timezone

from apps.accounts.models import Branch, CsPerm, Role, User
from apps.core.audit import record
from apps.core.notifications import notify_role, notify_users

from .models import (CARD_NO_MAX, RESUBMITTABLE, ActivationRequest, AT_BRANCH_FREE, Card, CardBlock, CardReason, CardStatus, CardUpload,
                     FollowUp, RejectReason, RequestStatus)


NAME_PATTERN = re.compile(r"^[^\W\d_]+(?: [^\W\d_]+)*$")     # letters only, words separated by single spaces


class CardError(Exception):
    """A rule prevented the action; the message is shown to the user as-is."""


def require_cs(user, flag: int) -> None:
    """Defence in depth: the views already check the right, the services check it again."""
    if not user.can_cs(flag):
        raise CardError("Your account is not allowed to do this. Ask an administrator for the right.")


def header_key(text) -> str:
    return "".join(ch for ch in str(text or "").lower() if ch.isalnum())


def _card_text(value) -> str:
    """A card number cell as text, never in scientific notation and never with a trailing ".0"."""
    if value is None:
        return ""
    if isinstance(value, float):
        value = int(Decimal(repr(value))) if value.is_integer() else value
    if isinstance(value, int):
        return str(value)
    return str(value).strip()


def read_card_numbers(data: bytes, filename: str) -> list[str]:
    """
    Extract ONLY the "Card No" column, in the exact order of the file (blank cells are skipped, nothing is sorted,
    filtered or changed). Raises CardError when the file has no "Card No" header.
    """
    extension = "." + filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if extension not in settings.ALLOWED_CARD_FILE_EXTENSIONS:
        raise CardError("Unsupported file type. Upload the card file as .xlsx or .csv.")
    if len(data) > settings.MAX_UPLOAD_MB * 1024 * 1024:
        raise CardError(f"The file is larger than {settings.MAX_UPLOAD_MB} MB.")
    if extension == ".xlsx":
        from openpyxl import load_workbook

        try:
            sheet = load_workbook(io.BytesIO(data), read_only=True, data_only=True).worksheets[0]
        except Exception as exc:  # noqa: BLE001
            raise CardError("The XLSX file could not be opened. Make sure it is a valid Excel file.") from exc
        rows = [list(r) for r in sheet.iter_rows(values_only=True)]
    else:
        for encoding in ("utf-8-sig", "cp1252"):
            try:
                text = data.decode(encoding)
                break
            except UnicodeDecodeError:
                continue
        else:
            raise CardError("The CSV file uses an unknown text encoding.")
        rows = [list(r) for r in csv.reader(io.StringIO(text))]
    for index, row in enumerate(rows[:30]):
        column = next((i for i, cell in enumerate(row) if header_key(cell) == "cardno"), None)
        if column is not None:
            numbers = [_card_text(r[column]) if column < len(r) else "" for r in rows[index + 1:]]
            numbers = [n for n in numbers if n]
            if not numbers:
                raise CardError('The "Card No" column is empty.')
            return numbers
    raise CardError('Could not find a column titled "Card No" in the first 30 rows of the file.')


@transaction.atomic
def import_card_file(user, data: bytes, filename: str) -> dict:
    """Store the card numbers of a file (status Not Sent, date = now). Numbers that are already stored are skipped."""
    require_cs(user, CsPerm.BATCH)
    numbers = read_card_numbers(data, filename)
    too_long = next((n for n in numbers if len(n) > CARD_NO_MAX), None)
    if too_long:
        raise CardError(f"Card number {too_long[:25]}... is longer than {CARD_NO_MAX} characters. Check the file.")
    existing = set(Card.objects.filter(card_no__in=numbers).values_list("card_no", flat=True))
    upload = CardUpload.objects.create(file_name=filename[:255], file_sha256=hashlib.sha256(data).hexdigest(), uploaded_by=user)
    now, seen, cards, skipped_known, skipped_repeat = timezone.now(), set(), [], 0, 0
    for position, number in enumerate(numbers):          # position = row order of the file
        if number in existing:
            skipped_known += 1
        elif number in seen:
            skipped_repeat += 1
        else:
            seen.add(number)
            cards.append(Card(upload=upload, position=position, card_no=number, created_at=now))
    Card.objects.bulk_create(cards)
    upload.card_count = len(cards)
    upload.save(update_fields=["card_count"])
    record("cards", upload.pk, "cards.uploaded", actor=user, file=filename, added=len(cards), already_stored=skipped_known, repeated=skipped_repeat)
    return {"upload": upload, "added": len(cards), "already_stored": skipped_known, "repeated": skipped_repeat}


@transaction.atomic
def deliver_cards(user, card_ids, branch: Branch) -> int:
    """CS: selected Not Sent cards -> Delivered to `branch`, date and time = now."""
    require_cs(user, CsPerm.BATCH)
    cards = list(Card.objects.select_for_update().filter(pk__in=card_ids, status=CardStatus.NOT_SENT))
    if not cards:
        raise CardError("Select at least one card that is still Not Sent.")
    now = timezone.now()
    Card.objects.filter(pk__in=[c.pk for c in cards]).update(status=CardStatus.DELIVERED, branch=branch, delivered_at=now, delivered_by=user)
    record("cards", 0, "cards.delivered", actor=user, branch=branch.code, count=len(cards))
    notify_users(User.objects.filter(role=Role.TELLER, branch=branch, is_active=True), "cards_delivered",
                 f"{len(cards)} card(s) were delivered to your branch ({branch.code}). Confirm them when they arrive.", "cards", 0)
    return len(cards)


@transaction.atomic
def accept_cards(user, card_ids) -> int:
    """Teller: Delivered cards of the teller's own branch -> Accepted, with the date of receipt."""
    cards = list(Card.objects.select_for_update().filter(pk__in=card_ids, status=CardStatus.DELIVERED, branch=user.branch))
    if not cards:
        raise CardError("Select at least one delivered card of your branch.")
    now = timezone.now()
    Card.objects.filter(pk__in=[c.pk for c in cards]).update(status=CardStatus.ACCEPTED, accepted_at=now, accepted_by=user)
    record("cards", 0, "cards.accepted", actor=user, branch=user.branch.code, count=len(cards))
    notify_role(Role.CS, "cards_accepted", f"Branch {user.branch.code} accepted {len(cards)} card(s).", "cards", 0)
    return len(cards)


# ---------------------------------------------------------------------------------------------------------------------
# Blocking cards (CS with the Block right)
# ---------------------------------------------------------------------------------------------------------------------
@transaction.atomic
def block_cards(user, card_ids, note: str) -> int:
    """
    CS (right "Block card"): cards that a branch still holds (Delivered / Accepted) -> Blocked, with a written reason.
    A card that already went to a customer or that is still at head office cannot be blocked here.
    """
    require_cs(user, CsPerm.BLOCK)
    note = " ".join((note or "").split())
    if len(note) < 3:
        raise CardError("Write the reason for blocking the card(s).")
    if len(note) > 255:
        raise CardError("The reason is too long (255 characters at most).")
    cards = list(Card.objects.select_for_update(of=("self",)).filter(pk__in=card_ids, status__in=AT_BRANCH_FREE).select_related("branch"))
    if not cards:
        raise CardError("Select at least one card that a branch still holds (Delivered or Accepted).")
    now = timezone.now()
    Card.objects.filter(pk__in=[c.pk for c in cards]).update(status=CardStatus.BLOCKED)
    CardBlock.objects.bulk_create([CardBlock(card=c, note=note, blocked_by=user, blocked_at=now) for c in cards])
    by_branch = defaultdict(int)
    for c in cards:
        by_branch[c.branch] += 1
    record("cards", 0, "cards.blocked", actor=user, count=len(cards), note=note, branches=sorted(b.code for b in by_branch))
    for branch, n in by_branch.items():
        notify_users(User.objects.filter(role=Role.TELLER, branch=branch, is_active=True), "cards_blocked",
                     f"{n} card(s) of your branch were blocked by Customer service: {note}", "cards", 0)
    return len(cards)


# ---------------------------------------------------------------------------------------------------------------------
# Dashboard: Total Undisbursed / Total Disbursed / Collected by Customer / Remaining / Block
# ---------------------------------------------------------------------------------------------------------------------
def dashboard(user) -> dict:
    """
    The numbers of the dashboard, counted live from the card table with ONE grouped query (nothing is stored, so the dashboard can
    never drift from the cards: a new batch, a delivery, a card request or a block changes the next read at once).

      Undisbursed  cards still at head office (Not Sent)                  - one figure for the whole company
      Disbursed    every card that was delivered to a branch (any later state)
      Collected    cards given to a customer
      Remaining    cards the branch still holds (Delivered + Accepted)
      Block        cards blocked by CS

    A branch that never received a card shows "-" (`empty`) instead of zeros, so a newly added branch starts as "Remaining Card: -".
    The branch of the signed-in Teller comes first (`mine`, shown in blue); everybody sees all branches.
    """
    per_branch = defaultdict(lambda: defaultdict(int))
    for branch_id, status, n in Card.objects.order_by().values_list("branch_id", "status").annotate(n=Count("pk")):
        per_branch[branch_id][status] += n

    def figures(counts) -> dict:
        disbursed = sum(counts.get(s, 0) for s in (CardStatus.DELIVERED, CardStatus.ACCEPTED, CardStatus.COLLECTED, CardStatus.BLOCKED))
        return {"disbursed": disbursed, "collected": counts.get(CardStatus.COLLECTED, 0),
                "remaining": counts.get(CardStatus.DELIVERED, 0) + counts.get(CardStatus.ACCEPTED, 0), "blocked": counts.get(CardStatus.BLOCKED, 0)}

    total = {"undisbursed": sum(c.get(CardStatus.NOT_SENT, 0) for c in per_branch.values()), "disbursed": 0, "collected": 0, "remaining": 0, "blocked": 0}
    branches = []
    for branch in Branch.objects.filter(is_active=True).order_by("name", "code"):
        item = {"id": branch.pk, "code": branch.code, "label": branch.name or branch.code, **figures(per_branch.get(branch.pk, {}))}
        item["empty"] = item["disbursed"] == 0
        item["mine"] = branch.pk == user.branch_id
        branches.append(item)
    for key in ("disbursed", "collected", "remaining", "blocked"):     # the totals include cards of a branch that was switched off
        total[key] = sum(figures(c)[key] for b, c in per_branch.items() if b is not None)
    branches.sort(key=lambda b: not b["mine"])          # the Teller's own branch first (stable: the rest stays in name order)
    return {"total": total, "branches": branches}


# ---------------------------------------------------------------------------------------------------------------------
# Card requests (customer data + PDF per customer)
# ---------------------------------------------------------------------------------------------------------------------
def validate_pdf(uploaded) -> None:
    name = uploaded.name or ""
    if not name.lower().endswith(".pdf"):
        raise CardError(f'"{name}" is not a PDF file.')
    if uploaded.size > settings.MAX_UPLOAD_MB * 1024 * 1024:
        raise CardError(f'"{name}" is larger than {settings.MAX_UPLOAD_MB} MB.')
    head = uploaded.read(5)
    uploaded.seek(0)
    if not head.startswith(b"%PDF"):
        raise CardError(f'"{name}" is not a valid PDF file.')


def suggest_cards(user, typed: str, limit: int = 8) -> list[str]:
    """
    The "search box" of the Teller form: card numbers the own branch can still use (Delivered / Accepted), most similar first.
    The number always starts with the fixed BIN; `typed` is what the Teller typed after it (digits only).
    Numbers that START with the typed digits come first, then numbers that merely CONTAIN them; both in the file order.
    """
    if user.role != Role.TELLER or not user.branch_id:
        return []
    typed = "".join(ch for ch in (typed or "") if ch.isdigit())
    bin_ = settings.CARD_BIN
    if typed.startswith(bin_):
        typed = typed[len(bin_):]
    free = Card.objects.filter(branch_id=user.branch_id, status__in=AT_BRANCH_FREE, card_no__startswith=bin_).order_by("upload_id", "position", "id")
    if not typed:
        return list(free.values_list("card_no", flat=True)[:limit])
    starts = list(free.filter(card_no__startswith=bin_ + typed).values_list("card_no", flat=True)[:limit])
    if len(starts) >= limit:
        return starts
    inside = free.annotate(rest=Substr("card_no", len(bin_) + 1)).filter(rest__contains=typed).exclude(card_no__startswith=bin_ + typed).values_list("card_no", flat=True)[:limit - len(starts)]
    return starts + list(inside)


def _clean(text) -> str:
    return " ".join(str(text or "").split())


def clean_request_data(data: dict) -> dict:
    """Check and normalise the customer fields of a card request (shared by the form and the service)."""
    name = _clean(data.get("customer_name")).upper()
    account = "".join(str(data.get("account_no") or "").split())
    id_no = "".join(str(data.get("id_no") or "").split()).upper()
    phone = "".join(str(data.get("phone") or "").split())
    if len(name) < 2:
        raise CardError("Enter the customer's name.")
    if len(name) > 100:
        raise CardError("The customer's name is too long (100 characters at most).")
    if not NAME_PATTERN.match(name):
        raise CardError("The name can only contain letters and spaces.")
    if not (account.isdigit() and 5 <= len(account) <= 20):
        raise CardError("The account number must be 5 to 20 digits.")
    if not (id_no.isalnum() and 5 <= len(id_no) <= 20):
        raise CardError("The ID number must be 5 to 20 letters or digits.")
    if not (phone.lstrip("+").isdigit() and 7 <= len(phone) <= 20):
        raise CardError("Enter the customer's phone number (7 to 20 digits).")
    try:
        reason = CardReason(int(data.get("reason")))
    except (TypeError, ValueError):
        raise CardError("Choose the reason for the card.") from None
    return {"customer_name": name, "account_no": account, "id_no": id_no, "phone": phone, "reason": reason}


@transaction.atomic
def submit_request(user, data: dict, uploaded, card_no: str) -> ActivationRequest:
    """
    Teller: one card request (status Sent) for one customer. The card must be one the own branch still holds. It becomes Collected by
    the customer at once (dashboard: Remaining - 1, Collected + 1). Nothing is saved when any check fails.
    """
    fields = clean_request_data(data)
    if uploaded is None:
        raise CardError("Choose the customer's PDF file.")
    validate_pdf(uploaded)
    card_no = "".join(str(card_no or "").split())
    if not (card_no.startswith(settings.CARD_BIN) and card_no.isdigit()):
        raise CardError(f"The card number must start with {settings.CARD_BIN} and contain digits only.")
    card = Card.objects.select_for_update().filter(card_no=card_no, branch_id=user.branch_id, status__in=AT_BRANCH_FREE).first()
    if card is None:
        raise CardError("This card is not available at your branch (not delivered to it, already used or blocked). Pick one from the suggestions.")
    now = timezone.now()
    if card.status == CardStatus.DELIVERED:           # the Teller holds it in the hand, so receipt is confirmed at the same moment
        card.accepted_at, card.accepted_by = now, user
    card.status, card.collected_at = CardStatus.COLLECTED, now
    card.save(update_fields=["status", "collected_at", "accepted_at", "accepted_by"])
    req = ActivationRequest.objects.create(branch=user.branch, submitted_by=user, card=card, file=uploaded, file_name=uploaded.name[:255],
                                           submitted_at=now, status_changed_at=now, **fields)
    record("activation", req.pk, "activation.submitted", actor=user, branch=user.branch.code, card=card.card_no, customer=fields["customer_name"])
    notify_role(Role.CS, "activation_sent", f"Branch {user.branch.code} sent a card request for {fields['customer_name']}.", "activation", req.pk)
    return req


ALLOWED_NEXT = {
    RequestStatus.SENT: {RequestStatus.ACTIVATION_REQUESTED, RequestStatus.REJECTED},
    RequestStatus.ACTIVATION_REQUESTED: {RequestStatus.ACTIVATED, RequestStatus.REJECTED},
}


@transaction.atomic
def update_request_status(user, request_id: int, status: int, follow_up=None, reject_reason=None) -> ActivationRequest:
    """
    CS (right "Process card activation"): Sent -> Activation Requested | Reject; Activation Requested -> Card Activated | Reject.
    A Reject needs a type (Incomplete Data / Return / Close) and a reason from the list.
    """
    require_cs(user, CsPerm.ACTIVATION)
    req = ActivationRequest.objects.select_for_update().get(pk=request_id)
    status = int(status)
    if status not in ALLOWED_NEXT.get(req.status, set()):
        raise CardError(f"The status cannot change from {req.get_status_display()} to that value.")
    now = timezone.now()
    if status == RequestStatus.REJECTED:
        try:
            follow_up, reject_reason = FollowUp(int(follow_up)), RejectReason(int(reject_reason))
        except (TypeError, ValueError):
            raise CardError("Choose the reject type and the reason.") from None
        req.follow_up, req.reject_reason = follow_up, reject_reason
    else:
        req.follow_up = req.reject_reason = None
        if status == RequestStatus.ACTIVATION_REQUESTED:
            req.mabes_at = now
        else:
            req.activated_at = now
    previous = req.status
    req.status, req.handled_by, req.status_changed_at = status, user, now
    req.save()
    record("activation", req.pk, "activation.status", actor=user, previous=previous, status=status,
           follow_up=req.follow_up, reason=req.reject_reason)
    text = f'Card request "{req.customer_name or req.file_name}": {req.get_status_display()}.'
    if status == RequestStatus.REJECTED:
        text += f" {req.get_follow_up_display()} - {req.get_reject_reason_display()}."
    notify_users(User.objects.filter(role=Role.TELLER, branch=req.branch, is_active=True), "activation_status", text, "activation", req.pk)
    return req


@transaction.atomic
def resubmit_request(user, request_id: int, uploaded) -> ActivationRequest:
    """Teller: a rejected request (Incomplete Data or Return) gets a corrected PDF and goes back to Sent. A Closed one cannot."""
    req = ActivationRequest.objects.select_for_update().get(pk=request_id, branch=user.branch)
    if req.status != RequestStatus.REJECTED or req.follow_up not in RESUBMITTABLE:
        raise CardError("Only a rejected request that is not Closed can be sent again.")
    validate_pdf(uploaded)
    old_file, old_name, old_note = req.file, req.file_name, req.note_text
    now = timezone.now()
    req.file, req.file_name = uploaded, uploaded.name[:255]
    req.status, req.status_changed_at, req.submitted_at = RequestStatus.SENT, now, now
    req.follow_up = req.reject_reason = None
    req.cs_note = ""
    req.resubmit_count += 1
    req.save()
    old_file.delete(save=False)
    record("activation", req.pk, "activation.resubmitted", actor=user, previous_file=old_name, previous_note=old_note)
    notify_role(Role.CS, "activation_sent", f'Branch {user.branch.code} sent "{req.customer_name or req.file_name}" again after a Reject.', "activation", req.pk)
    return req
