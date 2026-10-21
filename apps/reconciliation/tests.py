"""Tests for the bank readers, the matching rules, Finance review actions, retention and the Finance screens (revision 2)."""
from datetime import timedelta
from datetime import date
from decimal import Decimal
from io import BytesIO

from django.test import TestCase, override_settings
from django.urls import reverse
from django.utils import timezone
from openpyxl import load_workbook

from apps.core.models import AuditLog, TransactionDescription
from apps.core.normalize import normalize_reference
from apps.core.testing import (ALB_XLSX, PASSWORD, PBB_XLSX, SAMPLE_DIR, credit_row, make_world, qr_row, receipt, upload_generic,
                               upload_statement, xlsx_statement)

from .models import MatchStatus, ReconciliationMatch, StatementEntry, StatementImport
from .services import review
from .services.engine import run_matching
from .services.importer import start_import
from .services.parser import (parse_statement, rhb_qr_reference, rhb_sender_name, clean_bank_ref, split_alb_reference)
from .services.report import build_report
from .services.retention import purge_expired


def status(sub):
    sub.match.refresh_from_db()
    return sub.match.match_status


def upload_xlsx(user, bank, rows, name):
    import hashlib
    data = xlsx_statement(rows)
    imp = start_import(user=user, bank=bank, data=data, digest=hashlib.sha256(data).hexdigest(), file_name=name)
    imp.refresh_from_db()
    return imp


class ParserTests(TestCase):
    def test_rhb_helpers(self):
        self.assertEqual(rhb_sender_name(" ZARIEAM/ ALI/"), "ZARIEAM")
        self.assertEqual(rhb_qr_reference(" RHBQR049318/20260925TNGDMYNB030OQR85237862"), "QR85237862")
        self.assertEqual(rhb_qr_reference("/oce"), "")
        self.assertEqual(clean_bank_ref("'0000006932"), "0000006932")

    def test_alb_reference_2_is_name_slash_reference_third_part_ignored(self):
        self.assertEqual(split_alb_reference("YOSEPIN ISABELA/RPP260829309437120"), ("YOSEPIN ISABELA", "RPP260829309437120"))
        self.assertEqual(split_alb_reference("Javinci Berkat/RPP260828309298341/INARA GROUP SDN"), ("Javinci Berkat", "RPP260828309298341"))

    def test_rhb_real_sample(self):
        parsed = parse_statement((SAMPLE_DIR / "RHB_statement_sample.xlsx").read_bytes(), "a.xlsx", "rhb")
        first = parsed.rows[0]
        self.assertEqual((first.counterparty_name, first.reference, first.amount), ("ZARIEAM", "QR85417923", Decimal("500")))
        plain = next(r for r in parsed.rows if r.description == "RPP INWARD INST TRF")
        self.assertEqual(plain.reference, "0000006932")           # Bank Ref No., apostrophe removed

    def test_alb_real_sample_only_duitnow_cr_trf_ca_is_matchable(self):
        parsed = parse_statement((SAMPLE_DIR / "ALB_statement_sample.xlsx").read_bytes(), "a.xlsx", "alb")
        for row in parsed.rows:
            self.assertEqual(row.matchable, row.description == "DuitNow CR Trf CA")
        first = parsed.rows[0]
        self.assertEqual((first.counterparty_name, first.reference), ("YOSEPIN ISABELA", "RPP260829309437120"))
        self.assertTrue(parsed.open_item_only)

    def test_pbb_real_sample_only_duitnow_trsf_cr_is_matchable(self):
        parsed = parse_statement((SAMPLE_DIR / "PBB_statement_sample.xlsx").read_bytes(), "a.xlsx", "pbb")
        self.assertTrue(parsed.rows)
        for row in parsed.rows:
            self.assertEqual(row.matchable, row.description == "DUITNOW TRSF CR")


@override_settings(IMPORT_RUN_SYNC=True)
class MatchingTests(TestCase):
    def setUp(self):
        self.w = make_world()

    def test_exact_match_is_gfp(self):
        sub = receipt(self.w["teller"], self.w["rhb"], "150.00", "REF123456", sender="CUSTOMER ONE")
        upload_statement(self.w["finance"], [credit_row("REF123456", "150.00")])
        self.assertEqual(status(sub), MatchStatus.GFP)

    def test_name_is_compared_with_the_account_holder_not_the_sender(self):
        sub = receipt(self.w["teller"], self.w["rhb"], "150.00", "REF123456", sender="SOMEONE ELSE", holder="CUSTOMER ONE")
        upload_statement(self.w["finance"], [credit_row("REF123456", "150.00", name="CUSTOMER ONE")])
        self.assertEqual(status(sub), MatchStatus.GFP)

    def test_reference_is_not_a_matching_parameter_name_and_amount_are(self):
        sub = receipt(self.w["teller"], self.w["rhb"], "1000.00", "05831232", sender="KARYATI LESTARI", holder="KARYATI LESTARI")
        upload_statement(self.w["finance"], [credit_row("0000002895", "1000.00", name="LESTARI KARYATI", description="RPP INWARD INST TRF")])
        self.assertEqual(status(sub), MatchStatus.GFP)

    def test_name_difference_is_not_found_not_review(self):
        sub = receipt(self.w["teller"], self.w["rhb"], "150.00", "REF123456", sender="CUSTOMER ONE")
        upload_statement(self.w["finance"], [credit_row("REF123456", "150.00", name="TOTALLY DIFFERENT")])
        self.assertEqual(status(sub), MatchStatus.NOT_FOUND)

    def test_amount_difference_goes_to_review_with_reason(self):
        sub = receipt(self.w["teller"], self.w["rhb"], "150.00", "REF123456")
        upload_statement(self.w["finance"], [credit_row("OTHERREF99", "140.00")])
        sub.match.refresh_from_db()
        self.assertEqual(sub.match.match_status, MatchStatus.MANUAL_REVIEW)
        self.assertEqual(sub.match.review_reasons, ["amount"])

    def test_rhb_qr_rows_match_on_the_qr_id_not_the_bank_ref_no(self):
        sub = receipt(self.w["teller"], self.w["rhb"], "50.00", "QR84400381")
        upload_statement(self.w["finance"], [qr_row("QR84400381", "50.00", ref_no="0000009999")])
        self.assertEqual(status(sub), MatchStatus.GFP)

    def test_only_the_selected_bank_is_matched(self):
        rhb_sub = receipt(self.w["teller"], self.w["rhb"], "150.00", "REF123456")
        alb_sub = receipt(self.w["teller"], self.w["alb"], "150.00", "REF123456")
        upload_statement(self.w["finance"], [credit_row("REF123456", "150.00")])
        self.assertEqual(status(rhb_sub), MatchStatus.GFP)
        self.assertEqual(status(alb_sub), MatchStatus.PENDING)

    def test_same_name_and_amount_at_different_times_use_one_bank_row_each(self):
        first = receipt(self.w["teller"], self.w["rhb"], "150.00", "REF123456")
        second = receipt(self.w["teller"], self.w["rhb"], "150.00", "REF999999")
        upload_statement(self.w["finance"], [credit_row("B1", "150.00"), credit_row("B2", "150.00")])
        self.assertEqual((status(first), status(second)), (MatchStatus.GFP, MatchStatus.GFP))
        self.assertNotEqual(first.match.bank_statement_entry_id, second.match.bank_statement_entry_id)

    def test_one_bank_row_serves_one_receipt(self):
        first = receipt(self.w["teller"], self.w["rhb"], "150.00", "REF123456")
        second = receipt(self.w["teller2"], self.w["rhb"], "150.00", "REF123456X", sender="CUSTOMER ONE")
        upload_statement(self.w["finance"], [credit_row("REF123456", "150.00")])
        self.assertEqual(sorted([status(first), status(second)]).count(MatchStatus.GFP), 1)

    def test_alb_matches_duitnow_rows_and_keeps_other_rows_as_open_items(self):
        sub = receipt(self.w["teller"], self.w["alb"], "7000.00", "RPP260829000000001", sender="TESTER ONE")
        imp = upload_xlsx(self.w["finance"], self.w["alb"], ALB_XLSX, "alb.xlsx")
        self.assertEqual(status(sub), MatchStatus.GFP)
        cdm = StatementEntry.objects.get(description="CDM CA CSH DEP-W/O C")
        self.assertFalse(cdm.matchable)
        self.assertEqual(StatementEntry.objects.filter(matchable=False).count(), 2)
        self.assertEqual(imp.bank, self.w["alb"])

    def test_pbb_matches_only_duitnow_trsf_cr_rows(self):
        sub = receipt(self.w["teller"], self.w["pbb"], "22500.00", "338623", sender="TESTER TWO")
        other = receipt(self.w["teller"], self.w["pbb"], "1345.15", "000000", sender="FROM 3999XXXXXX")
        upload_xlsx(self.w["finance"], self.w["pbb"], PBB_XLSX, "pbb.xlsx")
        self.assertEqual(status(sub), MatchStatus.GFP)
        self.assertNotEqual(status(other), MatchStatus.GFP)

    def test_only_extracted_columns_are_stored(self):
        upload_statement(self.w["finance"], [credit_row("REF123456", "150.00")])
        names = {f.name for f in StatementEntry._meta.get_fields()}
        self.assertNotIn("raw_columns", names)
        self.assertNotIn("references", names)


@override_settings(IMPORT_RUN_SYNC=True)
class ReviewTests(TestCase):
    def setUp(self):
        self.w = make_world()
        self.sub = receipt(self.w["teller"], self.w["rhb"], "150.00", "REF123456")
        upload_statement(self.w["finance"], [credit_row("REF123456", "140.00")])
        self.match = ReconciliationMatch.objects.get(teller_submission=self.sub)

    def test_confirming_an_amount_difference_needs_a_note(self):
        with self.assertRaises(review.ReviewError):
            review.confirm_suggestion(self.match, self.w["finance"], "")
        review.confirm_suggestion(self.match, self.w["finance"], "customer paid less, agreed")
        self.assertEqual(status(self.sub), MatchStatus.GFP)

    def test_reject_sets_reject_keeps_the_compared_row_and_notifies(self):
        review.reject_receipt(self.match, self.w["finance"], "Amaun Salah")
        self.match.refresh_from_db()
        self.assertEqual(self.match.match_status, MatchStatus.REJECTED)
        self.assertEqual(self.match.teller_status, "rejected")
        self.assertIsNotNone(self.match.rejected_bank_statement_entry)
        self.assertTrue(self.sub.submitted_by.notifications.filter(type="rejected").exists())

    def test_reject_needs_a_listed_reason(self):
        with self.assertRaises(review.ReviewError):
            review.reject_receipt(self.match, self.w["finance"], "  ")

    def test_reject_reason_is_a_dropdown_value_and_lain_lain_keeps_a_remark(self):
        with self.assertRaises(review.ReviewError):
            review.reject_receipt(self.match, self.w["finance"], "free text reason")
        review.reject_receipt(self.match, self.w["finance"], "lain-lain", "bank slip unreadable")
        self.match.refresh_from_db()
        self.assertEqual(self.match.finance_note, "lain-lain: bank slip unreadable")

    def test_reject_through_the_screen_uses_the_dropdown(self):
        self.client.login(username="finance@t.test", password=PASSWORD)
        html = self.client.get(reverse("reconciliation:detail", args=[self.match.pk])).content.decode()
        for value in ("Nama Beda", "Amaun Salah", "lain-lain"):
            self.assertIn(f'<option value="{value}">', html)
        self.client.post(reverse("reconciliation:reject", args=[self.match.pk]), {"reject_reason": "Nama Beda"})
        self.match.refresh_from_db()
        self.assertEqual(self.match.finance_note, "Nama Beda")

    def test_rejected_receipt_is_not_rematched_until_corrected(self):
        review.reject_receipt(self.match, self.w["finance"], "Amaun Salah")
        upload_statement(self.w["finance"], [credit_row("REF123456", "150.00", name="CUSTOMER ONE")], name="b.csv", date="29/09/2026")
        self.assertEqual(status(self.sub), MatchStatus.REJECTED)

    def test_correction_goes_back_to_pending_and_rematches(self):
        review.reject_receipt(self.match, self.w["finance"], "Amaun Salah")
        data = {"sender_name": self.sub.sender_name, "account_holder_name": self.sub.account_holder_name, "relationship": self.sub.relationship,
                "bank": self.sub.bank, "amount": Decimal("140.00"), "bank_reference_no": "REF123456",
                "payment_method_code": self.sub.payment_method_code, "transaction_at": self.sub.transaction_at}
        review.correct_submission(self.match, self.w["teller"], data)
        self.assertEqual(status(self.sub), MatchStatus.PENDING)    # back to Pending, as the Teller sees it
        run_matching(match_ids=[self.match.pk], quiet=True, promote_only=True)
        self.assertEqual(status(self.sub), MatchStatus.GFP)      # re-checked at once against the stored bank row

    def test_unchanged_correction_is_refused(self):
        review.reject_receipt(self.match, self.w["finance"], "Amaun Salah")
        sub = self.sub
        data = {"sender_name": sub.sender_name, "account_holder_name": sub.account_holder_name, "relationship": sub.relationship, "bank": sub.bank,
                "amount": sub.amount, "bank_reference_no": sub.bank_reference_no, "payment_method_code": sub.payment_method_code,
                "transaction_at": sub.transaction_at}
        with self.assertRaises(review.ReviewError):
            review.correct_submission(self.match, self.w["teller"], data)


@override_settings(IMPORT_RUN_SYNC=True)
class RetentionTests(TestCase):
    def setUp(self):
        self.w = make_world()

    def test_bank_data_older_than_31_days_is_deleted_receipts_and_gfp_snapshot_stay(self):
        sub = receipt(self.w["teller"], self.w["rhb"], "150.00", "REF123456")
        imp = upload_statement(self.w["finance"], [credit_row("REF123456", "150.00"), credit_row("OTHER99", "10.00", name="X")])
        StatementImport.objects.filter(pk=imp.pk).update(imported_at=timezone.now() - timedelta(days=32))
        counts = purge_expired()
        self.assertEqual(counts["imports"], 1)
        self.assertEqual(StatementEntry.objects.count(), 0)
        sub.match.refresh_from_db()
        self.assertEqual(sub.match.match_status, MatchStatus.GFP)
        self.assertEqual(sub.match.bank_snapshot["reference"], "REF123456")

    def test_recent_data_is_kept(self):
        upload_statement(self.w["finance"], [credit_row("REF123456", "150.00")])
        StatementImport.objects.update(imported_at=timezone.now() - timedelta(days=30))
        self.assertEqual(purge_expired()["imports"], 0)

    def test_review_receipt_returns_to_pending_when_its_row_is_purged(self):
        sub = receipt(self.w["teller"], self.w["rhb"], "150.00", "REF123456")
        imp = upload_statement(self.w["finance"], [credit_row("REF123456", "140.00")])
        StatementImport.objects.filter(pk=imp.pk).update(imported_at=timezone.now() - timedelta(days=40))
        purge_expired()
        self.assertEqual(status(sub), MatchStatus.PENDING)


@override_settings(IMPORT_RUN_SYNC=True)
class FinanceScreenTests(TestCase):
    def setUp(self):
        self.w = make_world()
        self.sub = receipt(self.w["teller"], self.w["rhb"], "150.00", "REF123456")
        upload_statement(self.w["finance"], [credit_row("REF123456", "140.00"), credit_row("UNCLAIMED77", "75.00", name="STRANGER")])
        self.match = ReconciliationMatch.objects.get(teller_submission=self.sub)
        self.client.login(username="finance@t.test", password=PASSWORD)

    def test_review_table_has_no_gfp_or_reject_buttons_only_detail(self):
        html = self.client.get(reverse("reconciliation:queue") + "?tab=review").content.decode()
        self.assertIn(reverse("reconciliation:detail", args=[self.match.pk]), html)
        self.assertNotIn(reverse("reconciliation:confirm", args=[self.match.pk]), html)
        self.assertNotIn(reverse("reconciliation:reject", args=[self.match.pk]), html)

    def test_detail_has_note_gfp_and_reject_and_one_comparison(self):
        html = self.client.get(reverse("reconciliation:detail", args=[self.match.pk])).content.decode()
        self.assertIn(reverse("reconciliation:confirm", args=[self.match.pk]), html)
        self.assertIn(reverse("reconciliation:reject", args=[self.match.pk]), html)
        self.assertIn("is-diff", html)

    def test_after_reject_the_detail_still_compares_with_the_bank_row(self):
        self.client.post(reverse("reconciliation:reject", args=[self.match.pk]), {"reject_reason": "Amaun Salah"})
        self.match.refresh_from_db()
        self.assertEqual(self.match.match_status, MatchStatus.REJECTED)
        html = self.client.get(reverse("reconciliation:detail", args=[self.match.pk])).content.decode()
        self.assertIn("REF123456", html)
        self.assertIn("is-diff", html)

    def test_reject_tab_and_open_item_tab(self):
        self.client.post(reverse("reconciliation:reject", args=[self.match.pk]), {"reject_reason": "Amaun Salah"})
        self.assertIn("REF123456", self.client.get(reverse("reconciliation:queue") + "?tab=rejected").content.decode())
        html = self.client.get(reverse("reconciliation:queue") + "?tab=open_item").content.decode()
        self.assertIn("UNCLAIMED77", html)

    def test_description_filter_options_come_from_master_data(self):
        TransactionDescription.objects.create(name="My Custom Desc")
        html = self.client.get(reverse("reconciliation:queue") + "?tab=open_item").content.decode()
        self.assertIn("My Custom Desc", html)

    def test_teller_sees_reject_and_it_is_listed_first(self):
        other = receipt(self.w["teller"], self.w["rhb"], "20.00", "NEWER00001")
        review.reject_receipt(self.match, self.w["finance"], "Amaun Salah")
        self.client.logout()
        self.client.login(username="teller@t.test", password=PASSWORD)
        html = self.client.get(reverse("receipts:list")).content.decode()
        self.assertIn("Reject", html)
        self.assertLess(html.index("REF123456"), html.index("NEWER00001"))
        self.assertTrue(other.pk)


@override_settings(IMPORT_RUN_SYNC=True)
class ReportTests(TestCase):
    def test_report_has_description_filter_and_open_items(self):
        w = make_world()
        upload_statement(w["finance"], [credit_row("UNCLAIMED77", "75.00", name="STRANGER")])
        today = date(2026, 9, 30)   # the date of the uploaded statement
        data, _name = build_report("weekly", today, None, "RPP INWARD INST TRF CR")
        values = [str(c.value) for sheet in load_workbook(BytesIO(data)).worksheets for row in sheet.iter_rows() for c in row]
        self.assertIn("UNCLAIMED77", values)
        header = [c.value for c in load_workbook(BytesIO(data))["Transactions"][1]]
        self.assertNotIn("Receipt #", header)
        self.assertNotIn("Name", header)
        self.assertEqual(header.count("Account holder"), 1)
        data, _name = build_report("weekly", today, None, "SI CR")
        values = [str(c.value) for sheet in load_workbook(BytesIO(data)).worksheets for row in sheet.iter_rows() for c in row]
        self.assertNotIn("UNCLAIMED77", values)
