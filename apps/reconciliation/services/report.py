"""
Excel report for Finance: bank transactions of one week or one month with their status.

GFP        The bank row was matched to a Teller receipt.
Open Item  The bank row has no matching Teller receipt yet.

The period is applied to the transaction date of the bank row when the bank file
has a date column, otherwise to the statement date of the upload.

Accepted input date format: MM/DD/YYYY (e.g. 10/12/2026).
"""

import calendar
import io
import re
from datetime import date, datetime, timedelta

from django.db.models.functions import Coalesce
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from apps.core.models import Bank

from ..models import StatementEntry


HEADERS = [
    "Bank",
    "Date",
    "Reference",
    "Account holder",
    "Description",
    "Direction",
    "Amount (RM)",
    "Status",
    "Branch",
    "Reported by",
    "Teller reference",
    "Receipt time",
]

WIDTHS = [12, 12, 22, 26, 34, 10, 14, 12, 8, 20, 20, 18]
RECEIPT_TIME_COL = HEADERS.index("Receipt time") + 1


def clean_teller_text(value):
    """
    Normalise teller-entered text: UPPERCASE, no symbols (e.g. "/", "(", ")", "-"),
    single spaces. Letters and digits are kept. None / empty stays None.
    Example: "Tina Teller (KJ)" -> "TINA TELLER KJ"
    """
    if value is None:
        return None
    text = re.sub(r"[^A-Za-z0-9\s]", " ", str(value))
    text = re.sub(r"\s+", " ", text).strip().upper()
    return text or None


def parse_anchor_date(anchor: date | str) -> date:
    """
    Convert the input date into a date object.

    Accepted string format:
        MM/DD/YYYY, e.g. 10/12/2026

    Existing date objects are accepted without conversion.
    """
    if isinstance(anchor, datetime):
        return anchor.date()

    if isinstance(anchor, date):
        return anchor

    if isinstance(anchor, str):
        try:
            return datetime.strptime(anchor.strip(), "%m/%d/%Y").date()
        except ValueError as exc:
            raise ValueError(
                "Invalid date format. Use MM/DD/YYYY, for example 10/12/2026."
            ) from exc

    raise TypeError(
        "The anchor date must be a date object or a string in MM/DD/YYYY format."
    )


def period_bounds(
    period: str,
    anchor: date | str,
) -> tuple[date, date]:
    """
    Weekly = Monday to Sunday containing the anchor date.
    Monthly = the calendar month containing the anchor date.

    The anchor date accepts MM/DD/YYYY strings or date objects.
    """
    anchor = parse_anchor_date(anchor)

    if period == "weekly":
        start = anchor - timedelta(days=anchor.weekday())
        return start, start + timedelta(days=6)

    if period == "monthly":
        start = anchor.replace(day=1)
        end = anchor.replace(
            day=calendar.monthrange(anchor.year, anchor.month)[1]
        )
        return start, end

    raise ValueError("Invalid period. Use 'weekly' or 'monthly'.")


def report_filename(period: str, start: date, end: date) -> str:
    return (
        f"Reconciliation_Report_{period.capitalize()}_"
        f"{start:%Y-%m-%d}_to_{end:%Y-%m-%d}.xlsx"
    )


def transactions(
    start: date,
    end: date,
    bank_id: int | None = None,
    description: str = "",
):
    qs = (
        StatementEntry.objects
        .annotate(
            report_date=Coalesce(
                "transaction_date",
                "statement_import__statement_date",
            )
        )
        .filter(
            report_date__gte=start,
            report_date__lte=end,
        )
        .select_related(
            "bank",
            "statement_import",
            "match",
            "match__teller_submission",
            "match__teller_submission__branch",
            "match__teller_submission__submitted_by",
        )
        .order_by(
            "report_date",
            "bank__name",
            "source_row",
            "id",
        )
    )

    qs = qs.filter(bank_id=bank_id) if bank_id else qs
    return qs.filter(description__iexact=description) if description else qs


def build_report(
    period: str,
    anchor: date | str,
    bank_id: int | None = None,
    description: str = "",
) -> tuple[bytes, str]:
    """
    Generate the Excel report.

    Args:
        period: 'weekly' or 'monthly'.
        anchor: Date object or date string in MM/DD/YYYY format.
        bank_id: Optional bank ID to filter the report.

    Returns:
        A tuple containing the Excel file bytes and its filename.
    """
    # Convert MM/DD/YYYY input into a date object and calculate the period.
    anchor = parse_anchor_date(anchor)
    start, end = period_bounds(period, anchor)

    font = Font(name="Arial", size=10)
    bold = Font(name="Arial", size=10, bold=True)
    head_font = Font(
        name="Arial",
        size=10,
        bold=True,
        color="FFFFFF",
    )
    head_fill = PatternFill("solid", fgColor="2554E8")
    thin = Side(style="thin", color="CBD5E1")
    box = Border(
        left=thin,
        right=thin,
        top=thin,
        bottom=thin,
    )
    green = PatternFill("solid", fgColor="E3F6EC")
    amber = PatternFill("solid", fgColor="FFF3D6")

    wb = Workbook()
    ws = wb.active
    ws.title = "Transactions"
    ws.append(HEADERS)

    for col, width in enumerate(WIDTHS, start=1):
        cell = ws.cell(row=1, column=col)
        cell.font = head_font
        cell.fill = head_fill
        cell.border = box
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True,
        )
        ws.column_dimensions[get_column_letter(col)].width = width

    ws.row_dimensions[1].height = 24

    row_no = 1
    # bank name -> [gfp_count, open_count, gfp_amount, open_amount]
    totals = {}

    for entry in transactions(start, end, bank_id, description):
        match = getattr(entry, "match", None)
        sub = match.teller_submission if match else None
        status = "GFP" if match else "Open Item"

        t = totals.setdefault(entry.bank.name, [0, 0, 0, 0])
        idx = 0 if match else 1
        t[idx] += 1
        t[idx + 2] += entry.amount or 0

        ws.append([
            entry.bank.name,
            entry.report_date,
            entry.display_reference,
            entry.counterparty_name,
            entry.description,
            entry.direction,
            entry.amount,
            status,
            clean_teller_text(sub.branch.code) if sub else None,
            clean_teller_text(
                sub.submitted_by.name if sub and sub.submitted_by else None
            ),
            clean_teller_text(sub.bank_reference_no) if sub else None,
            (
                timezone.make_naive(timezone.localtime(sub.transaction_at))
                if sub
                else None
            ),
        ])

        row_no += 1

        for col in range(1, len(HEADERS) + 1):
            cell = ws.cell(row=row_no, column=col)
            cell.font = font
            cell.border = box

        ws.cell(row=row_no, column=2).number_format = "dd mmm yyyy"
        ws.cell(row=row_no, column=7).number_format = "#,##0.00"
        ws.cell(row=row_no, column=RECEIPT_TIME_COL).number_format = "dd mmm yyyy hh:mm"

        ws.cell(row=row_no, column=8).fill = green if match else amber

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = (
        f"A1:{get_column_letter(len(HEADERS))}{max(row_no, 2)}"
    )

    # Summary: per-bank totals computed above.
    sm = wb.create_sheet("Summary")

    sm["A1"] = "Reconciliation report"
    sm["A1"].font = Font(name="Arial", size=14, bold=True)

    sm["A2"] = "Period"
    sm["B2"] = (
        f"{period.capitalize()}: "
        f"{start:%d %b %Y} to {end:%d %b %Y}"
    )

    sm["A3"] = "Generated"
    sm["B3"] = timezone.localtime().strftime("%d %b %Y %H:%M")

    for r in (2, 3):
        sm[f"A{r}"].font = bold
        sm[f"B{r}"].font = font

    heads = [
        "Bank",
        "GFP (count)",
        "Open Item (count)",
        "Total (count)",
        "GFP (RM)",
        "Open Item (RM)",
        "Total (RM)",
    ]

    for i, h in enumerate(heads, start=1):
        c = sm.cell(row=5, column=i, value=h)
        c.font = head_font
        c.fill = head_fill
        c.border = box
        c.alignment = Alignment(
            horizontal="center",
            wrap_text=True,
        )
        sm.column_dimensions[get_column_letter(i)].width = (
            18 if i > 1 else 16
        )

    banks = (
        [Bank.objects.get(pk=bank_id)]
        if bank_id
        else list(
            Bank.objects.exclude(statement_format="")
            .order_by("name")
        )
    )

    last = 5

    # Static values (not formulas): formulas written by openpyxl have no cached
    # result, so Protected View, previewers, pandas etc. show them as empty.
    for r, bank in enumerate(banks, start=6):
        gfp_n, open_n, gfp_amt, open_amt = totals.get(bank.name, [0, 0, 0, 0])
        sm.cell(row=r, column=1, value=bank.name)
        sm.cell(row=r, column=2, value=gfp_n)
        sm.cell(row=r, column=3, value=open_n)
        sm.cell(row=r, column=4, value=gfp_n + open_n)
        sm.cell(row=r, column=5, value=gfp_amt)
        sm.cell(row=r, column=6, value=open_amt)
        sm.cell(row=r, column=7, value=gfp_amt + open_amt)
        last = r

    total_row = last + 1

    sm.cell(row=total_row, column=1, value="Total")

    for col in range(2, 8):
        sm.cell(
            row=total_row,
            column=col,
            value=sum(sm.cell(row=r, column=col).value for r in range(6, last + 1)),
        )

    for r in range(6, total_row + 1):
        for col in range(1, 8):
            c = sm.cell(row=r, column=col)
            c.font = bold if r == total_row else font
            c.border = box

            if col in (5, 6, 7):
                c.number_format = "#,##0.00"

    buffer = io.BytesIO()
    wb.save(buffer)

    return buffer.getvalue(), report_filename(period, start, end)
