"""

Bank statement reader (CSV or XLSX) for the four bank categories.

Only the columns that matching needs are extracted (see bank_profiles.py); every other column of the file is ignored.

* RHB keeps its fixed layout: row 1 = "Balance of Account <number> / ... as at <date>" (read for metadata, BR-3),

  row 3 = "Current Account" (ignored, BR-4), header on row 4 (BR-5).

    name       "Sender`s Name/ Beneficiary's Name": the part BEFORE the first "/" (" ZARIEAM/ ALI/" -> "ZARIEAM").

    amount     Credit or Debit (a row has only one).

    reference  Description "DUITNOW QR POS CR": the "QR<digits>" at the end of "Reference 1/ Recipient`s Reference"

               (" RHBQR049318/20260925TNGDMYNB030OQR85237862" -> "QR85237862").

               Any other Description: "Bank Ref No." ("'0000006932" -> "0000006932").

* ALB: name and reference come from "Reference 2" ("NAME/REFERENCE" or "NAME/REFERENCE/RECEIVING ACCOUNT NAME", the third

  value is dropped). Only "DuitNow CR Trf CA" rows are matched; all other rows are stored as Open Items (matchable=False).

* PBB: name = "Reference 2", reference = "Cheque No/Ref No", amount = Debit Amount / Credit Amount. Only "DUITNOW TRSF CR"

  rows are matched ("SI CR", "TSFR FUND CR-ATM/EFT", ... are Open Items; they share reference numbers, so they are never matched).

* MBB: generic reader (reference = "Transaction Ref", no name).

"""
import csv
import hashlib
import io
import re
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal
from django.conf import settings
from apps.core.normalize import normalize_reference, parse_amount
from .bank_profiles import (FIELD_LABELS, GENERIC_ALIASES, PROFILES, BankProfile, description_key, detect_bank, has_required, key,
                            map_headers, squash)

class StatementParseError(Exception):
    """Raised with a user-readable message when the file cannot be used."""

class Row1UnreadableError(StatementParseError):
    """RHB row 1 has no readable account number / date. The upload form then asks for them manually."""

class BankMismatchError(StatementParseError):
    """The file looks like another bank's statement than the one selected."""
    def __init__(self, detected: str, selected: str):
        self.detected, self.selected = detected, selected
        super().__init__(f"This file looks like an {PROFILES[detected].label} statement, but {PROFILES[selected].label} is selected. "
                         f"Select {PROFILES[detected].label} and upload it again.")

DATE_FORMATS = ("%d/%m/%Y", "%d-%m-%Y", "%d %b %Y", "%d %B %Y", "%d-%b-%Y", "%d-%b-%y", "%Y-%m-%d", "%d/%m/%y", "%d.%m.%Y", "%d %b %y", "%Y/%m/%d")

SUMMARY_LINE = re.compile(r"total|opening balance|closing balance|balance b/?f|balance c/?f|brought forward|carried forward|baki", re.I)
# RHB: Description that carries a QR payment. Its reference is the "QR<digits>" of "Reference 1/ Recipient`s Reference".

RHB_QR_DESCRIPTION = "DUITNOW QR POS CR"
# An identifier segment: optional letters, then at least 6 digits, no spaces ("RPP260829309437120").
_ID_SEGMENT = re.compile(r"^[A-Za-z]{0,6}\d{6,}[A-Za-z0-9]*$")
# RHB: the QR id is the LAST "QR<digits>" of the text. The first "QR" belongs to the merchant code "RHBQR049318".
_RHB_QR_AT_END = re.compile(r"QR\d+$", re.I)

def alb_is_credit_transfer(description: str) -> bool:
    return PROFILES["alb"].is_matchable(description)

def split_alb_reference(text: str) -> tuple[str, str]:
    """

    ALB "Reference 2" -> (name, reference). The first value is the name, the second the reference, a third value is ignored.

    "YOSEPIN ISABELA/RPP260829309437120"                               -> ("YOSEPIN ISABELA", "RPP260829309437120")

    "Javinci Berkat Kreat/RPP260828309298341/INARA GROUP SDN. BHD"     -> ("Javinci Berkat Kreat", "RPP260828309298341")

    "KUMAR A/L RAJAN/RPP260828309298341"                               -> ("KUMAR A/L RAJAN", "RPP260828309298341")

    A name may itself contain a "/" ("A/L"), so the reference is the first segment that looks like an identifier and

    everything before it is the name. Without such a segment the plain rule applies: first value = name, second = reference.

    """
    parts = [p.strip() for p in (text or "").split("/")]
    for index, part in enumerate(parts):
        if index and part and _ID_SEGMENT.match(part):
            return squash("/".join(parts[:index])), part
    parts = [p for p in parts if p]
    if len(parts) >= 2:
        return squash(parts[0]), parts[1]
    return squash(parts[0]) if parts else "", ""

def rhb_sender_name(text: str) -> str:
    """" ZARIEAM/ ALI/" -> "ZARIEAM": only the Sender`s Name, i.e. the part before the first "/"."""
    return squash((text or "").split("/", 1)[0])

def rhb_qr_reference(text: str) -> str:
    """" RHBQR049318/20260925TNGDMYNB030OQR85237862" -> "QR85237862". Empty string when the text does not end in a QR id."""
    match = _RHB_QR_AT_END.search(squash(text))
    return match.group(0).upper() if match else ""

def clean_bank_ref(text: str) -> str:
    """"'0000006932" -> "0000006932" (Excel text marker removed)."""
    return (text or "").strip().lstrip("'`\u2019 ").strip()

@dataclass
class ParsedRow:
    source_row: int
    description: str
    debit: Decimal | None
    credit: Decimal | None
    amount: Decimal
    counterparty_name: str        # compared with the Teller's ACCOUNT HOLDER
    reference: str                # as shown ("QR85237862")
    bank_ref_no: str              # unique key of the row (duplicate protection)
    reference_normalized: str
    transaction_date: date | None
    matchable: bool = True        # False: Open Item straight away

@dataclass
class ParsedStatement:
    bank_key: str
    account_number: str
    account_label: str
    statement_date: date
    file_format: str
    rows: list[ParsedRow] = field(default_factory=list)
    skipped: list[str] = field(default_factory=list)   # human-readable reasons for rows we could not use
    duplicates_in_file: int = 0
    open_item_only: int = 0                            # rows of a transaction type that is not reconciled (stored as Open Items)
    header_row: int = 0
    columns: list[str] = field(default_factory=list)

def _clean_excel_literal(value: object) -> str:
    """Unwrap CSV values exported by Excel as formulas, e.g. ="2184.25"."""
    text = "" if value is None else str(value).strip()
    if not text:
        return ""
    match = re.fullmatch(r'=\s*"((?:[^"]|"")*)"', text)
    if match:
        return match.group(1).replace('""', '"').strip()
    return text

def _cell(value) -> str:
    """Excel returns numbers as float/int/datetime; turn everything into clean text."""
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d %H:%M:%S") if (value.hour or value.minute or value.second) else value.strftime("%Y-%m-%d")
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return _clean_excel_literal(value)

def read_rows(data: bytes, extension: str) -> list[list[str]]:
    if extension == ".xlsx":
        from openpyxl import load_workbook
        try:
            workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        except Exception as exc:
            raise StatementParseError("The XLSX file could not be opened. Make sure it is a valid Excel file "
                                      "downloaded from the bank and not password protected.") from exc
        sheet = workbook.worksheets[0]
        return [[_cell(c) for c in row] for row in sheet.iter_rows(values_only=True)]
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            text = data.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    else:
        raise StatementParseError("The CSV file uses an unknown text encoding. Re-download it from the bank website.")
    return [[_cell(c) for c in row] for row in csv.reader(io.StringIO(text))]

def parse_date(text: str) -> date | None:
    text = (text or "").strip()
    if not text:
        return None
    if re.match(r"^\d{4}-\d{2}-\d{2}[ T]", text):
        text = text[:10]
    for fmt in DATE_FORMATS:
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    # "29/08/2026 07:30:49" (ALB) and "07-oct-2026 MY (UTC+08:00)" (MBB): use the date part only
    match = re.search(r"\d{1,2}[/\-. ][A-Za-z]{3,9}[/\-. ]\d{2,4}|\d{1,2}[/\-. ]\d{1,2}[/\-. ]\d{2,4}", text)
    if match and match.group(0) != text:
        return parse_date(match.group(0))
    return None

def _rhb_metadata(row1: list[str]) -> tuple[str, str, date | None]:
    """Row 1: 'Balance of Account 1234567890 / ... as at 30/09/2026' (BR-3)."""
    text = " ".join(c for c in row1 if c)
    match = re.search(r"Account\s+([A-Za-z0-9\-]+)", text, re.IGNORECASE)
    if match:
        account = match.group(1)
    else:
        digits = re.search(r"\d{6,}", text)
        account = digits.group(0) if digits else ""
    when = None
    match = re.search(r"as at\s+(.+)$", text, re.IGNORECASE)
    if match:
        when = parse_date(match.group(1).strip().rstrip("."))
    return account, text[:100], when

def _generic_metadata(rows: list[list[str]]) -> tuple[str, str, date | None]:
    """Account number and statement date from the lines above the header, when the bank prints them."""
    text = " ".join(" ".join(r) for r in rows)
    match = re.search(r"(?:account\s*(?:no\.?|number)?|a/c\s*(?:no\.?)?)\s*[:\-]?\s*([0-9][0-9\- ]{5,19}[0-9])", text, re.IGNORECASE)
    account = match.group(1).replace(" ", "") if match else ""
    when = None
    # a printed label followed by a date in the next cell ("To Date | 2026-10-07", "Date | 2026-10-07"), best label first
    found = {}
    for row in rows:
        cells = [c for c in row if c]
        for i, cell in enumerate(cells[:-1]):
            found.setdefault(key(cell), parse_date(cells[i + 1]))
    for label in ("statementdate", "todate", "asat", "asof", "enddate", "date", "generationdate"):
        if found.get(label):
            when = found[label]
            break
    if not when:
        match = re.search(r"(?:as at|as of|statement date|period to|to)\s*[:\-]?\s*(\d{1,2}[/\-. ][A-Za-z0-9]{1,9}[/\-. ]\d{2,4})", text, re.IGNORECASE)
        if match:
            when = parse_date(match.group(1))
    return account, text[:100].strip(), when

def _find_header(rows: list[list[str]], profile: BankProfile) -> int:
    for index, row in enumerate(rows[:60]):
        if has_required(map_headers(row, profile), profile):
            return index
    raise StatementParseError(f"This file does not look like a {profile.label} statement. Check the selected bank and choose the bank statement file.")


def _find_column_index(header: list[str], *labels: str) -> int | None:
    """Find a column by its displayed header, independently of bank_profiles aliases."""
    wanted = {key(label) for label in labels}
    return next((index for index, title in enumerate(header) if key(title) in wanted), None)


def _row_value(row: list[str], index: int | None) -> str:
    """Safely read a cell by header index."""
    return squash(row[index]) if index is not None and index < len(row) else ""


def _split_amount_header(rows: list[list[str]], header_index: int, mapping: dict) -> tuple[int, dict]:
    """ALB prints ONE 'Amount' cell with 'Debit' and 'Credit' in the row below it. Returns (first data row index, extra titles)."""
    if "amount" in mapping and "debit" not in mapping and "credit" not in mapping and header_index + 1 < len(rows):
        sub = [key(c) for c in rows[header_index + 1]]
        debit = next((i for i, k in enumerate(sub) if k in GENERIC_ALIASES["debit"]), None)
        credit = next((i for i, k in enumerate(sub) if k in GENERIC_ALIASES["credit"]), None)
        if debit is not None and credit is not None:
            mapping.pop("amount")
            mapping["debit"], mapping["credit"] = debit, credit
            return header_index + 2, {debit: "Debit", credit: "Credit"}
    return header_index + 1, {}

def _synthetic_key(values: list[str], seen: dict) -> str:
    """Banks without a unique reference column: fingerprint the whole row. The n-th identical row gets '#n'."""
    digest = hashlib.sha1("|".join(key(v) for v in values).encode()).hexdigest()[:10]
    seen[digest] = seen.get(digest, 0) + 1
    return f"ROW-{digest}-{seen[digest]}"

def parse_statement(data: bytes, filename: str, bank: str = "rhb", fallback_account: str = "",
                    fallback_date: date | None = None) -> ParsedStatement:
    extension = "." + filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if extension not in settings.ALLOWED_STATEMENT_EXTENSIONS:
        raise StatementParseError("Unsupported file type. Upload the bank statement as a .csv or .xlsx file.")
    if bank not in PROFILES:
        raise StatementParseError("This bank has no statement reader. Ask the administrator to set the statement format in Master data.")
    profile = PROFILES[bank]
    rows = read_rows(data, extension)
    if not any(any(r) for r in rows):
        raise StatementParseError("The file is empty.")
    detected = detect_bank(rows)
    if detected and detected != bank:
        raise BankMismatchError(detected, bank)
    # ---- header row and metadata
    if profile.fixed_layout:
        header_index = profile.header_row - 1
        if len(rows) <= header_index:
            raise StatementParseError(f"The file has only {len(rows)} row(s). A {profile.label} statement needs the header on row "
                                      f"{profile.header_row} followed by transactions.")
        account, label, when = _rhb_metadata(rows[0])
        account, when = account or fallback_account.strip(), when or fallback_date
        if not account or not when:
            raise Row1UnreadableError(
                "Row 1 has no readable account number or statement date (expected: 'Balance of Account 1234567890 / "
                "... as at 30/09/2026'). Enter both below and upload again.")
        mapping = map_headers(rows[header_index], profile)
        if not has_required(mapping, profile):
            missing = [FIELD_LABELS[f] for f in ("description", "debit", "credit", "bank_ref_no") if f not in mapping]
            raise StatementParseError(
                f"Row {profile.header_row} must contain the column headers, but these are missing: {', '.join(missing)}. "
                f"Check that this is an {profile.label} statement with rows 1-3 untouched and the header on row {profile.header_row}.")
    else:
        header_index = _find_header(rows, profile)
        mapping = map_headers(rows[header_index], profile)
        account, label, when = _generic_metadata(rows[:header_index])
        account = account or fallback_account.strip()
        when = when or fallback_date
    data_start, renamed = (header_index + 1, {}) if profile.fixed_layout else _split_amount_header(rows, header_index, mapping)
    header = rows[header_index]
    # PBB puts the depositor/account-holder name in Reference 2 and the receipt
    # reference in Cheque No/Ref No. Read these by actual header labels so an
    # incomplete bank_profiles alias cannot turn the fields into empty strings.
    pbb_name_index = _find_column_index(header, "Reference 2") if bank == "pbb" else None
    pbb_reference_index = (
        _find_column_index(header, "Cheque No/Ref No", "Cheque No / Ref No", "Cheque No/Ref. No.")
        if bank == "pbb" else None
    )
    titles = []                                  # unique, non-empty column titles (shown by `inspect_statement`)
    for i, title in enumerate(header):
        name = title.strip() or f"Column {i + 1}"
        titles.append(name if name not in titles else f"{name} ({i + 1})")
    parsed = ParsedStatement(bank_key=bank, account_number=account, account_label=label, statement_date=when or date.today(),
                             file_format=extension.lstrip("."), header_row=header_index + 1, columns=titles)
    for index, name in renamed.items():
        if index < len(titles):
            titles[index] = name
    # every column that may hold the date (MBB leaves "Transaction Date" empty and fills "Posting Date")
    date_indexes = [i for i, title in enumerate(header) if key(title) in profile.aliases.get("date", ())]
    def get(row: list[str], name: str) -> str:
        idx = mapping.get(name)
        return row[idx].strip() if idx is not None and idx < len(row) else ""
    seen_refs: set[str] = set()
    seen_fingerprints: dict = {}
    for offset, row in enumerate(rows[data_start:], start=data_start + 1):
        if not any(row):
            continue  # blank line
        if any(re.match(r"\s*total\b", cell, re.I) for cell in row if cell):
            continue  # "Total Amount" footer line
        description = squash(get(row, "description"))
        if SUMMARY_LINE.search(description) and not get(row, "bank_ref_no"):
            continue  # totals / opening / closing balance lines are not transactions
        try:
            debit, credit = parse_amount(get(row, "debit")), parse_amount(get(row, "credit"))
            single = parse_amount(get(row, "amount"))
        except ValueError as exc:
            if sum(1 for c in row if c) >= 3:
                parsed.skipped.append(f"Row {offset}: {exc}.")
            continue
        if single is not None and debit is None and credit is None:   # one signed column
            debit, credit = (abs(single), None) if single < 0 else (None, single)
        amount = credit if credit else debit                          # a row has only one of the two filled
        if not amount:
            if sum(1 for c in row if c) >= 3:
                parsed.skipped.append(f"Row {offset}: neither Debit nor Credit has an amount.")
            continue
        matchable = profile.is_matchable(description)
        name, reference, real_ref = "", "", clean_bank_ref(get(row, "bank_ref_no"))
        # ---- row key (duplicate protection)
        if profile.fixed_layout:
            if not real_ref:
                parsed.skipped.append(f"Row {offset}: Bank Ref No. is empty.")
                continue
            if real_ref in seen_refs:
                parsed.duplicates_in_file += 1
                continue
            seen_refs.add(real_ref)
            row_key = real_ref
        else:
            row_key = _synthetic_key(row, seen_fingerprints)    # these banks have no unique reference column
        # ---- name (= account holder) and reference: only the columns of the bank's rule
        if bank == "rhb":
            name = rhb_sender_name(get(row, "name"))
            if description_key(description) == RHB_QR_DESCRIPTION:
                reference = rhb_qr_reference(get(row, "reference_1")) or real_ref
            else:
                reference = real_ref
        elif bank == "alb":
            cell = get(row, "reference_2")
            if matchable:
                name, reference = split_alb_reference(cell)
            else:
                reference = squash(cell)          # shown in the Open Item list / report only
        elif bank == "pbb":
            # PBB: name from Reference 2 (e.g. "RIMA AMELIA"); reference from
            # Cheque No/Ref No (e.g. "535408"). Use actual header indexes first.
            name = _row_value(row, pbb_name_index) or squash(get(row, "reference_2"))
            reference = clean_bank_ref(_row_value(row, pbb_reference_index)) or real_ref
        else:                                     # mbb
            name, reference = squash(get(row, "name")), squash(get(row, "reference_1"))
        reference = reference.strip()
        if not matchable:
            parsed.open_item_only += 1
        normalized = normalize_reference(reference)
        if normalized and set(normalized) == {"0"}:
            normalized = ""                       # "000000" (PBB standing instructions) is not a reference
        parsed.rows.append(ParsedRow(
            source_row=offset, description=description[:100], debit=debit, credit=credit, amount=amount,
            counterparty_name=name[:150], reference=reference[:150], bank_ref_no=row_key[:50], reference_normalized=normalized[:150],
            transaction_date=next((d for d in (parse_date(row[i]) for i in date_indexes if i < len(row)) if d), None),
            matchable=matchable and bool(normalized)))
    if not parsed.rows:
        raise StatementParseError("No usable transactions were found below the header row.")
    if not when:   # banks that do not print a statement date: the latest transaction date is the best answer
        dates = [r.transaction_date for r in parsed.rows if r.transaction_date]
        parsed.statement_date = max(dates) if dates else (fallback_date or date.today())
    return parsed
